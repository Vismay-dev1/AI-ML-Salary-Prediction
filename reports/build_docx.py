"""
Builds reports/AI_ML_Salary_Prediction_Project_Guide.docx  (beginner-friendly written guide).
Run build_charts.py first.
"""
import json, os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "reports", "assets"); CH = os.path.join(A, "charts")
M = json.load(open(os.path.join(A, "metrics.json")))
OUT = os.path.join(ROOT, "reports", "AI_ML_Salary_Prediction_Project_Guide.docx")
R = M["results"]; CV = M["cv"]; IMP = M["importance_rf"]
money = lambda v: f"${v:,.0f}"
k = lambda v: f"${v/1000:,.0f}k"
INDIGO = RGBColor(0x63, 0x66, 0xF1); NAVY = RGBColor(0x0F, 0x17, 0x2A); SLATE = RGBColor(0x64, 0x74, 0x8B)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(8); st.paragraph_format.line_spacing = 1.2
for name, size, color in [("Heading 1", 22, INDIGO), ("Heading 2", 15, NAVY), ("Heading 3", 12.5, NAVY)]:
    h = doc.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    h.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12); h.paragraph_format.space_after = Pt(6)

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)

def H1(t): doc.add_heading(t, level=1)
def H2(t): doc.add_heading(t, level=2)
def H3(t): doc.add_heading(t, level=3)
def P(t, bold_lead=None, italic=False, size=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " "); r.bold = True
    r = p.add_run(t); r.italic = italic
    if size: r.font.size = Pt(size)
    return p
def BUL(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " "); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)
        p.paragraph_format.space_after = Pt(4)
def NUM(items):
    for it in items:
        p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(4)
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " "); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)

def CALLOUT(title, text, color="E0E7FF", title_color=INDIGO):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, color)
    p = c.paragraphs[0]; r = p.add_run(title); r.bold = True; r.font.color.rgb = title_color
    for line in ([text] if isinstance(text, str) else text):
        q = c.add_paragraph(); q.add_run(line); q.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def CODE(code, caption=None):
    if caption: P(caption, italic=True, size=10)
    t = doc.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, "F1F5F9")
    lines = code.strip("\n").split("\n")
    p = c.paragraphs[0]
    for i, line in enumerate(lines):
        r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if i < len(lines) - 1: r.add_break()
    p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def TABLE(rows, widths=None, header_color="0F172A", zebra=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c_, val in enumerate(row):
            cell = t.cell(r, c_); cell.text = ""
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(val)); run.font.size = Pt(10)
            if r == 0:
                run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(cell, header_color)
            elif zebra and r % 2 == 0: shade(cell, "F8FAFC")
            if widths: cell.width = Inches(widths[c_])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

FIG = [0]
def FIGURE(name, caption, width=6.3):
    FIG[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(CH, name), width=Inches(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figure {FIG[0]}. {caption}"); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = SLATE

# ======================================================= COVER
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(20)
p.add_run().add_picture(os.path.join(A, "hero.png"), width=Inches(6.4))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AI/ML Job Salary Prediction"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A complete, beginner-friendly guide to what this project does, how the data is used, and what the results mean")
r.font.size = Pt(13); r.font.color.rgb = SLATE
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Project by Vismay Vinod  •  Repository: github.com/Vismay-dev1/AI-ML-Salary-Prediction"); r.font.size = Pt(10); r.font.color.rgb = SLATE
doc.add_page_break()

H1("Contents")
toc = ["1. What is this project?", "2. Key words you will meet", "3. The big picture: seven steps",
       "4. The dataset", "5. Exploring the data", "6. Preparing the data for the computer",
       "7. Training and testing models", "8. Results: which model is best?", "9. How accurate are the predictions?",
       "10. What the model pays attention to", "11. Making your own salary prediction",
       "12. Limitations and honest notes", "13. How to run the project yourself", "14. Repository map",
       "15. Learning path and next steps", "16. Quick self-test", "Appendix: notebook cell guide"]
for t in toc: P(t)
CALLOUT("How to read this guide",
        ["Every technical word is explained the first time it appears, and there is a glossary in Section 2.",
         "Charts are shown as numbered figures. Code is shown in grey boxes, with a line-by-line explanation.",
         "All numbers come from re-running the project's notebook logic on the CSV file, so they match the code."])
doc.add_page_break()

# ======================================================= 1
H1("1. What is this project?")
P("Imagine you are about to apply for an AI or machine-learning job. You want to know: \"What salary should I expect?\" "
  "Salaries depend on many things at once: how many years you have worked, your degree, how many skills you know, "
  "which city you work in, and so on. Working out the pattern by hand is hard.")
P("This project tries to solve that problem with machine learning. Machine learning means giving a computer many past examples "
  "so it can find patterns on its own. Here the examples are 500 AI/ML job records. The computer studies them, "
  "and then it can estimate a salary for a job it has never seen before.")
H3("In one sentence")
CALLOUT("The project in one sentence", "It uses 500 past AI/ML job records to train four regression models, compares how well they predict salary (USD), and shows how to predict the salary for a new job profile.")
H3("What is inside the repository?")
TABLE([["File", "What it is", "Why it matters"],
       ["ai_ml_job_analysis.csv", "The dataset: 500 rows, 11 columns", "The raw information the models learn from"],
       ["AI_ML_Salary_Prediction.ipynb", "A Google Colab notebook with all the code and explanations", "Where the whole analysis happens, step by step"],
       ["README.md", "Short project description and how to run it", "The first thing to read on GitHub"],
       ["LICENSE", "MIT licence", "Tells you that you may reuse the code"]], widths=[2.2, 2.4, 2.8])
P("A notebook (.ipynb) is a document that mixes text, code and charts. You run one block of code at a time and see the result right below it. "
  "Google Colab is a free website that runs notebooks in your browser, so you do not need to install anything.")

# ======================================================= 2
H1("2. Key words you will meet")
P("Do not worry if these look unfamiliar. You will see each of them explained in context later in the guide.")
TABLE([["Word", "Plain-English meaning"],
       ["Dataset", "A table of examples. Here: 500 rows, one per job."],
       ["Feature (input)", "A piece of information the model uses to make a guess, e.g. Experience (Years)."],
       ["Target (output)", "The thing we want to predict. Here: Salary (USD)."],
       ["Model", "A recipe the computer learns from data to make predictions."],
       ["Regression", "Predicting a number (like salary) rather than a category (like 'yes/no')."],
       ["Training", "The model studies examples where the answer is already known."],
       ["Test set", "Examples kept aside, so we can check the model on jobs it has never seen."],
       ["Encoding", "Converting words (like 'Masters') into numbers the computer can use."],
       ["Correlation", "A score from -1 to +1 showing how two numbers move together."],
       ["MAE", "Mean Absolute Error: the average size of a mistake, in dollars. Lower is better."],
       ["RMSE", "Root Mean Squared Error: like MAE but punishes big mistakes more. Lower is better, in dollars."],
       ["R² (R-squared)", "How much of the salary variation the model explains. 1.0 is perfect; 0 means no better than guessing the average."],
       ["Feature importance", "How much the model relies on each input when it makes a decision."],
       ["Cross-validation", "Repeating the test on several different splits of the data to check the result is not luck."],
       ["Outlier", "A value far away from the rest, e.g. a very high salary."]], widths=[1.8, 4.9])

# ======================================================= 3
H1("3. The big picture: seven steps")
P("Every machine-learning project follows roughly the same path. The notebook follows these seven steps in order:")
FIGURE("c16_workflow.png", "The seven steps of the project, from raw data to a prediction.", width=6.5)
NUM([("Load the data:", "read the CSV file into a table the computer can work with (pandas)."),
     ("Explore:", "draw charts and calculate averages to see what is going on."),
     ("Prepare:", "convert text into numbers and drop the unhelpful ID column."),
     ("Split:", "keep 20% of the jobs aside as a test. The other 80% is for learning."),
     ("Train:", "fit four different models on the learning part."),
     ("Evaluate:", "score each model on the hidden test jobs with MAE, RMSE and R²."),
     ("Predict:", "use the best model to estimate the salary for a new job profile.")])
CALLOUT("Analogy: baking a cake", "Gather ingredients (data), taste them (explore), wash and chop (prepare), keep some batter aside for a taste test (split), bake several recipes (train), taste-test them (evaluate), then serve the best one (predict).", color="FEF3C7", title_color=RGBColor(0x92, 0x40, 0x0E))

# ======================================================= 4
H1("4. The dataset")
H2("4.1 What is in the file?")
P(f"The file ai_ml_job_analysis.csv has {M['rows']} rows and {M['cols']} columns. A CSV (comma-separated values) file is a plain text table: "
  "each line is a row, and commas separate the columns. You can open it in Excel or Google Sheets.")
TABLE([["Column", "What it means", "Example from row 3"],
       ["ID", "Row number. Not useful for predicting salary, so it is removed.", "3"],
       ["Job Title", "The name of the role (15 different titles).", "AI Consultant"],
       ["Experience (Years)", "Years of professional experience (0 to 20).", "13.0"],
       ["Education", "Highest degree: Bachelors, Masters or PhD.", "Masters"],
       ["Skill Count", "How many technical skills are listed (3 to 25).", "16"],
       ["Industry", "The sector the job is in (10 sectors).", "Technology"],
       ["Company", "The employer (20 companies).", "Cognizant"],
       ["Location", "City or country (15 values, e.g. 'London, UK').", "London, UK"],
       ["Remote Work", "Yes, No or Hybrid.", "No"],
       ["Certification", "A professional certificate, or 'None'.", "IBM Data Science Cert"],
       ["Salary (USD)", "The salary in US dollars. This is the TARGET we predict.", "143,200"]], widths=[1.6, 3.4, 1.7])

H2("4.2 Quick numbers")
TABLE([["Measure", "Value", "What it tells us"],
       ["Number of job records", f"{M['rows']}", "A small dataset, good for learning"],
       ["Average salary", money(M["salary_mean"]), "The typical salary if you add everything up and divide"],
       ["Median salary", money(M["salary_median"]), "The middle salary when all are sorted: half earn more, half less"],
       ["Lowest / highest salary", f"{money(M['salary_min'])} / {money(M['salary_max'])}", "The full range"],
       ["Different job titles", f"{M['n_titles']}", "Roles from Research Intern to AI Solutions Architect"],
       ["Different companies", f"{M['n_companies']}", "Big tech, consultancies and a startup"],
       ["Duplicate rows", f"{M['duplicates']}", "No copy-pasted records"]], widths=[1.9, 1.6, 3.2])
P("Why is the average higher than the median? Because a few very large salaries pull the average upward. "
  "This is called right skew.")

H2("4.3 How are the salaries spread out?")
FIGURE("c01_salary_distribution.png", "Histogram of all 500 salaries, with the mean (red) and median (amber) marked.")
P("A histogram counts how many jobs fall into each salary range. The tallest bars show where most jobs sit. "
  "The dashed lines show that the mean sits slightly above the median, which is the signature of a few high-paying outliers.")

H2("4.4 Checking the data quality")
P("Before trusting any model, a good analyst checks the data. Here is what the checks found:")
TABLE([["Check", "Result", "What we do about it"],
       ["Missing values", f"{M['missing_cert']} rows have Certification = 'None'. Pandas reads the text 'None' as missing.", "Keep them as their own category. Their average salary ($89k) is nearly the same as those with a certificate ($89k)."],
       ["Duplicate rows / IDs", "None found", "Nothing to fix."],
       ["Strange labels", f"{M['interns_over_5y']} 'Research Intern' rows report more than 5 years of experience.", "Keep them, but treat the patterns with care. The data may be partly synthetic."],
       ["Inconsistent places", f"Some companies appear in cities they do not normally operate in (e.g. Infosys in Toronto). {M['remote_location_rows']} rows use 'Remote' as a location.", "Noted as a limitation. The model learns from the data as it is."]], widths=[1.5, 2.7, 2.5])
FIGURE("c09_certification.png", "Left: average salary by certification. Right: 33% of rows have no certification ('None').", width=6.4)

# ======================================================= 5
H1("5. Exploring the data")
P("Exploratory data analysis (EDA) means looking at the data before modelling. The question is: which columns seem to move with salary? "
  "Each subsection below shows one chart and what it means.")
H2("5.1 Location: the biggest driver")
FIGURE("c03_location.png", "Average salary by location, coloured by region.", width=6.2)
rm = M["region_means"]; rc = M["region_counts"]
P(f"Location has the strongest effect in this data. Jobs in India average about {k(rm['India'])}, "
  f"while jobs in the USA average about {k(rm['USA'])} and Europe/Canada/Singapore about {k(rm['Europe / Canada / Singapore'])}. "
  f"Jobs marked 'Remote' average {k(rm['Remote'])}. Those are very large gaps, and the model confirms it later in Section 10.")

H2("5.2 Experience: more years, higher pay")
FIGURE("c06_experience_scatter.png", "Each dot is one job. Colour shows education. The red line is the overall trend.", width=6.0)
P(f"Each dot is a job. The overall trend slopes upward, and the correlation between experience and salary is {M['corr_exp_salary']:.2f}. "
  "A correlation of +1 would mean a perfect straight line, 0 means no link. 0.36 is a moderate positive link.")
FIGURE("c07_experience_bands.png", "Average salary grows with each band of experience.", width=5.6)
P(f"Grouping experience into bands makes the pattern easier to see: from {k(M['exp_bands']['0-3 yrs']['mean'])} for 0 to 3 years, "
  f"up to {k(M['exp_bands']['15-20 yrs']['mean'])} for 15 to 20 years.")

H2("5.3 Education")
FIGURE("c02_education.png", "Salary spread by education level. The label shows the median (middle value).", width=5.8)
P(f"Average salary rises with education: Bachelors {k(M['edu_means']['Bachelors'])}, Masters {k(M['edu_means']['Masters'])}, "
  f"PhD {k(M['edu_means']['PhD'])}. Notice that the boxes overlap a lot, meaning a PhD does not guarantee a higher salary than a Masters. "
  "Education is a useful but not decisive factor.")
H3("What is a box plot?")
P("A box plot shows the middle half of the data as a box, the median as a line in the box, and the spread as 'whiskers'. "
  "A wider box means more variation.")

H2("5.4 Remote work")
FIGURE("c08_remote.png", "Average salary for each remote-work type.", width=5.2)
P("Remote work makes almost no difference to average salary in this data. Remote, hybrid and on-site jobs all sit around $87k-$91k.")

H2("5.5 Job titles")
FIGURE("c04_job_title.png", "Average salary by job title. Red bars are below $60k.", width=6.0)
tm = M["title_means"]; tt = list(tm.keys())
P(f"{tt[0]} tops the list at about {k(tm[tt[0]])}, while Research Intern is lowest at about {k(tm[tt[-1]])}. "
  "Many middle-ranked titles are close to each other in average pay, so the title alone is not a strong guide.")

H2("5.6 Industry")
FIGURE("c05_industry.png", "Average salary by industry.", width=5.6)
P("Industry shows a smaller spread than location. Healthcare and Telecommunications lead; Retail is lowest on average, though its median is also low.")

H2("5.7 Numbers that move together: correlation")
FIGURE("c10_correlation.png", "Correlation between the three numeric columns. Red = move together, blue = move in opposite directions.", width=4.6)
TABLE([["Pair", "Correlation", "Meaning"],
       ["Experience and salary", f"{M['corr_exp_salary']:+.2f}", "Moderate positive link"],
       ["Skill count and salary", f"{M['corr_skill_salary']:+.2f}", "Weak positive link"],
       ["Experience and skill count", f"{M['corr_exp_skill']:+.2f}", "Practically no link"]], widths=[2.4, 1.3, 3.0])
CALLOUT("Beginner tip: correlation is not causation", "Two things moving together does not prove one causes the other. Experience may correlate with salary for many reasons (seniority, negotiation, promotions). The model shows association, not cause.")

# ======================================================= 6
H1("6. Preparing the data for the computer")
P("Computers work with numbers, not words. The notebook converts each text column into numbers in a specific way. "
  "There are three different strategies, depending on the kind of column.")
TABLE([["Step", "Column(s)", "What is done", "Why"],
       ["1", "ID", "Dropped", "It is just a row number, so it carries no information about salary."],
       ["2", "Education", "Bachelors = 0, Masters = 1, PhD = 2", "These have a natural order (more study = higher number)."],
       ["3", "Remote Work", "One-hot encoding: three new 0/1 columns (Remote_Yes, Remote_No, Remote_Hybrid)", "Yes, No and Hybrid have no natural order, so we avoid implying one is 'bigger'."],
       ["4", "Job Title, Industry, Company, Location, Certification", "Label encoding: each distinct value gets a number", "Models need numbers. Note: the numbers are assigned in alphabetical order, which carries no real meaning (see Section 12)."],
       ["5", "All features (linear model only)", "Standard scaling: each column is shifted and resized to a similar range", "Helps the straight-line model treat all inputs fairly."]], widths=[0.5, 1.6, 2.5, 2.1])
H3("Example: what the encoding looks like in code")
CODE('''data["Education"] = data["Education"].map({"Bachelors": 0, "Masters": 1, "PhD": 2})
data = pd.get_dummies(data, columns=["Remote Work"], prefix="Remote")
for col in ["Job Title", "Industry", "Company", "Location", "Certification"]:
    le = LabelEncoder()
    data[col] = le.fit_transform(data[col].astype(str))''',
     caption="The core of the preprocessing (simplified from notebook Section 3).")
BUL([("data[\"Education\"].map({...})", "looks up each word in the dictionary and replaces it with its number."),
     ("pd.get_dummies(...)", "creates one 0/1 column per remote option. A 'Yes' row gets Remote_Yes = 1 and the others 0."),
     ("LabelEncoder()", "scans a column, finds all unique values, and gives each one a number. 'fit_transform' does both steps at once."),
     ("astype(str)", "makes sure every value is text first, so missing values do not cause errors.")])
P("After preparation the table has 11 input columns (features) and one target column (Salary). The model sees these 11 inputs for every job.")
CALLOUT("The notebook's features", ["Job Title, Experience (Years), Education, Skill Count, Industry, Company, Location, Certification, Remote_Hybrid, Remote_No, Remote_Yes."], color="E0E7FF")

# ======================================================= 7
H1("7. Training and testing models")
H2("7.1 The train / test split")
P("We cannot grade a student with the same questions they studied. In the same way, we cannot judge a model on the jobs it learned from. "
  "So the 500 jobs are split into two groups:")
TABLE([["Group", "Size", "Used for"],
       ["Training set", f"{M['train_size']} jobs (80%)", "The model learns patterns from these"],
       ["Test set", f"{M['test_size']} jobs (20%)", "Hidden during training; used only to grade the model"]], widths=[1.8, 1.8, 3.1])
CODE('''X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)''',
     caption="The split, from the notebook (Section 4).")
BUL([("X", "the 11 input columns."), ("y", "the salary column (the target)."),
     ("test_size=0.2", "20% of rows go to the test set."),
     ("random_state=42", "a fixed 'seed' so the random shuffle is the same every time. Anyone running the code gets the same numbers.")])

H2("7.2 The four models, with analogies")
TABLE([["Model", "How it works (simple version)", "Analogy"],
       ["Linear Regression", "Draws the best straight line (or flat plane in many dimensions) through the data.", "A ruler"],
       ["Decision Tree", "Asks a chain of yes/no questions ('Is experience above 10?') until it reaches an answer. Max depth 8.", "A flowchart"],
       ["Random Forest", "Builds 300 separate trees, each on a random part of the data, and averages their answers. Max depth 12.", "Asking 300 experts and taking the average"],
       ["Gradient Boosting", "Builds 300 small trees one after another; each new tree corrects the mistakes of the ones before. Depth 3, learning rate 0.05.", "Correcting your own homework over and over"]], widths=[1.6, 3.4, 1.7])
CODE('''lr = LinearRegression().fit(X_train_scaled, y_train)
dt = DecisionTreeRegressor(max_depth=8, random_state=42).fit(X_train, y_train)
rf = RandomForestRegressor(n_estimators=300, max_depth=12, random_state=42, n_jobs=-1).fit(X_train, y_train)
gb = GradientBoostingRegressor(n_estimators=300, learning_rate=0.05, max_depth=3, random_state=42).fit(X_train, y_train)''',
     caption="Training all four models (simplified from notebook Section 5).")
BUL([(".fit(X, y)", "is the learning step: the model studies the inputs and their correct answers."),
     (".predict(X)", "is the guessing step: the trained model gives salary estimates."),
     ("Linear Regression uses the scaled inputs", "because straight-line models behave better when all inputs are on a similar scale. Tree models do not need this."),
     ("n_estimators=300", "is the number of trees; more trees make the forest more stable but slower.")])

# ======================================================= 8
H1("8. Results: which model is best?")
P("Three scores are used to compare the models on the 100 test jobs. Here is how to read each one:")
BUL([("MAE:", "on average, how far off the guess is, in dollars. Lower is better."),
     ("RMSE:", "similar to MAE but big misses count more. Lower is better."),
     ("R²:", "the share of salary variation the model explains. Closer to 1 is better.")])
rows = [["Model", "R²", "MAE (avg error)", "RMSE"]]
for n in ["Linear Regression", "Decision Tree", "Random Forest", "Gradient Boosting"]:
    rows.append([n, f"{R[n]['R2']:.3f}", money(R[n]["MAE"]), money(R[n]["RMSE"])])
TABLE(rows, widths=[2.4, 1.2, 1.9, 1.6])
P(f"A guess that always says the average salary would have an average error of about {money(M['baseline_mae'])}. "
  f"Linear Regression (MAE {money(R['Linear Regression']['MAE'])}) is barely better than that baseline, showing that a straight line cannot capture this data well. "
  f"Gradient Boosting is the clear winner, with an average error of {money(R['Gradient Boosting']['MAE'])} and R² of {R['Gradient Boosting']['R2']:.3f}.")
FIGURE("c11_model_comparison.png", "Comparison of the four models on R² (left), RMSE (middle) and MAE (right).", width=6.5)
H2("8.1 Is the result just luck? Cross-validation")
P("A single split could be lucky. Cross-validation repeats the process five times, each time using a different fifth of the data as the test set, "
  "and averages the scores. A good model should score well on every fold.")
FIGURE("c12_cross_validation.png", "Average R² over 5 folds, with the spread (error bars).", width=5.6)
P(f"Gradient Boosting averages R² of {CV['Gradient Boosting']['mean']:.3f} (± {CV['Gradient Boosting']['std']:.3f}), "
  f"Random Forest {CV['Random Forest']['mean']:.3f}, Decision Tree {CV['Decision Tree']['mean']:.3f}, and Linear Regression {CV['Linear Regression']['mean']:.3f}. "
  "The ranking is the same on every split, so the conclusion is robust.")

# ======================================================= 9
H1("9. How accurate are the predictions?")
FIGURE("c13_pred_vs_actual.png", "Predicted vs actual salary for the 100 test jobs (Gradient Boosting). Dots on the dashed line are perfect.", width=5.2)
P("Read this chart like a bullseye: the closer the dots are to the dashed diagonal, the better. The shaded band marks ±$20,000. "
  f"About {M['resid_within_20k']*100:.0f}% of predictions land inside that band.")
FIGURE("c14_residuals.png", "Errors (actual minus predicted). Most errors sit near zero.", width=5.6)
P(f"The error (residual) is the actual salary minus the predicted salary. The errors cluster around zero, with an average of about "
  f"-${abs(M['resid_mean']):,.0f}, which means the model is not consistently too high or too low.")

# ======================================================= 10
H1("10. What the model pays attention to")
P("Feature importance tells us how much the Random Forest relied on each input when making its decisions. "
  "The numbers add up to 100%.")
FIGURE("c15_feature_importance.png", "Feature importance from the Random Forest.", width=6.2)
TABLE([["Feature", "Importance", "Plain-English reading"],
       ["Location", f"{IMP['Location']*100:.1f}%", "The most powerful signal. Where the job is strongly shapes the salary."],
       ["Experience (Years)", f"{IMP['Experience (Years)']*100:.1f}%", "Second most important; more experience, more pay."],
       ["Job Title", f"{IMP['Job Title']*100:.1f}%", "The role has a noticeable effect."],
       ["Skill Count", f"{IMP['Skill Count']*100:.1f}%", "Matters somewhat, but less than the notebook's conclusion implies."],
       ["Remote_* columns (total)", f"{sum(v for kk, v in IMP.items() if kk.startswith('Remote'))*100:.1f}%", "Very little influence."]], widths=[2.0, 1.2, 3.5])
P("Gradient Boosting gives a very similar picture: Location is still first (about 62%), then Experience (about 19%). "
  "Note that the notebook's written conclusion says experience and skill count are the strongest numeric predictors; the chart confirms experience, but skill count is only fourth.")

# ======================================================= 11
H1("11. Making your own salary prediction")
P("Section 9 of the notebook lets you describe a candidate and get an estimate. The example profile in the notebook is:")
CODE("\n".join(f'    "{kk}": {vv!r},' for kk, vv in M["example_profile"].items()).join(["new_profile = {\n", "\n}"]),
     caption="The example profile from the notebook.")
P("Walking through the steps:")
NUM(["The profile is put into a one-row table.",
     "Education is converted to its number (Masters = 1).",
     "Remote Work = 'Yes' is turned into the Remote_Yes column.",
     "Each text column is converted to the number learned during training.",
     "Any missing one-hot columns are filled with 0, and the columns are put in the same order as training.",
     "The trained model returns the predicted salary."])
CALLOUT("Important catch: an unseen city silently changes the answer",
        [f"The notebook uses 'Bengaluru, India', but the dataset spells it 'Bangalore, India'. The label encoder has never seen 'Bengaluru', so the code falls back to the first known location alphabetically: '{M['example_fallback_location']}'.",
         f"As written, the notebook predicts Random Forest {money(M['example_pred_rf'])} and Gradient Boosting {money(M['example_pred_gb'])}.",
         f"With the correct spelling 'Bangalore, India', the predictions are Random Forest {money(M['example_pred_rf_fixed'])} and Gradient Boosting {money(M['example_pred_gb_fixed'])}.",
         "Lesson: always check that category names match the training data exactly, and warn the user when a value is unknown."],
        color="FEE2E2", title_color=RGBColor(0x9F, 0x12, 0x39))

# ======================================================= 12
H1("12. Limitations and honest notes")
BUL([("Small dataset.", "500 rows is enough for learning, but it is small for real-world salary prediction. A few unusual jobs can shift the results."),
     ("Data origin is unclear.", "The repository does not say where the data comes from. Some labels look implausible (e.g. many interns with 5+ years of experience), so the data may be partly synthetic. Treat results as a learning exercise, not a market report."),
     ("Label encoding imposes arbitrary order.", "Location, Company, Job Title and similar columns get numbers in alphabetical order. Tree models cope, but Linear Regression can be misled. The notebook calls this 'Frequency/Target-safe encoding', but it is plain LabelEncoder."),
     ("Unseen values fall back silently.", "The Bengaluru example shows how this can distort results (Section 11)."),
     ("One split for the written conclusion.", "Notebook metrics come from one 80/20 split. Cross-validation (Section 8.1) supports the ranking, but the notebook itself does not run it."),
     ("No saved outputs.", "The notebook was saved without its outputs, so the numbers in this guide were recomputed from the same code. They match the logic exactly."),
     ("Correlation is not causation.", "The model finds patterns, not reasons."),
     ("Mismatched columns.", "Some company/location pairs and 'Remote' used as a location are not realistic; a cleaner dataset would separate office location from employee location.")])

# ======================================================= 13
H1("13. How to run the project yourself")
H2("Option A: Google Colab (easiest)")
NUM(["Go to colab.research.google.com and choose File → Upload notebook. Select AI_ML_Salary_Prediction.ipynb.",
     "Run the first code cell. When asked, upload ai_ml_job_analysis.csv.",
     "Choose Runtime → Run all, or run each cell with Shift + Enter.",
     "Scroll through the charts and results."])
H2("Option B: on your own computer")
P("You need Python 3 and a few libraries. Open a terminal in the project folder and run:")
CODE("pip install pandas numpy matplotlib seaborn scikit-learn jupyter\njupyter notebook AI_ML_Salary_Prediction.ipynb", caption="Setup commands.")
P("To regenerate the charts, metrics and this guide, the reports folder contains three scripts:")
CODE("pip install python-pptx python-docx pillow\npython reports/build_charts.py     # charts + metrics.json\npython reports/build_pptx.py      # the interactive PowerPoint\npython reports/build_docx.py      # this Word guide", caption="Rebuilding the deliverables.")

# ======================================================= 14
H1("14. Repository map")
TABLE([["Path", "Purpose"],
       ["ai_ml_job_analysis.csv", "The dataset"],
       ["AI_ML_Salary_Prediction.ipynb", "The original analysis notebook"],
       ["README.md / LICENSE", "Project description and MIT licence"],
       ["reports/build_charts.py", "Re-runs the analysis and draws every chart"],
       ["reports/build_pptx.py", "Builds the interactive presentation"],
       ["reports/build_docx.py", "Builds this Word guide"],
       ["reports/assets/charts/", "All chart images used in the guide and deck"],
       ["reports/assets/metrics.json", "Every number quoted in the deck and guide"]], widths=[2.8, 3.9])

# ======================================================= 15
H1("15. Learning path and next steps")
P("If you are learning machine learning, this project is a good first end-to-end example. Here is a sensible path:")
NUM([("Understand the data:", "open the CSV in a spreadsheet and sort, filter and count values yourself."),
     ("Reproduce the charts:", "run the notebook and change one chart (e.g. colour or bins)."),
     ("Try a new feature:", "add a feature (e.g. 'Years Since Degree') and see if R² improves."),
     ("Tune the models:", "use GridSearchCV to search for better settings automatically."),
     ("Try stronger libraries:", "XGBoost or LightGBM are popular gradient-boosting tools."),
     ("Fix the encoding:", "replace LabelEncoder with one-hot or target encoding and compare results."),
     ("Build an app:", "wrap the trained model in a Streamlit web page where anyone can enter a profile.")])

# ======================================================= 16
H1("16. Quick self-test")
TABLE([["Question", "Answer"],
       ["1. What is the target column?", "Salary (USD)"],
       ["2. Why do we keep 100 jobs hidden?", "To test the model on jobs it has never seen, so we do not reward memorisation."],
       ["3. What does R² = 0.935 mean?", "The model explains about 93.5% of the salary variation on the test jobs."],
       ["4. Which feature matters most to the Random Forest?", f"Location (about {IMP['Location']*100:.0f}%)."],
       ["5. Why is Remote work less useful than expected?", "Its importance is under 1% in the model, and average salaries are similar across remote types."],
       ["6. Why must category names match exactly?", "An unseen name is silently mapped to another location, changing the prediction."]], widths=[3.0, 3.7])

# ======================================================= Appendix
H1("Appendix: notebook cell guide")
P("This table maps each part of the notebook to what it does, in plain language.")
TABLE([["Section", "What it does"],
       ["1. Load the dataset", "Checks for the CSV and reads it into a table called df."],
       ["2. Exploratory data analysis", "Shows the shape, data types, missing values, salary histogram, correlation heatmap, box plots by category and an experience-vs-salary scatter."],
       ["3. Data preprocessing", "Drops ID, encodes Education (ordinal), Remote Work (one-hot), and five text columns (label encoding)."],
       ["4. Train/test split", "Splits 80/20 with random_state = 42 and scales inputs for Linear Regression."],
       ["5. Model training", "Trains Linear Regression, Decision Tree, Random Forest and Gradient Boosting; prints MAE, RMSE and R²."],
       ["6. Model comparison", "Shows the results table and bar charts for R² and RMSE."],
       ["7. Best model: predicted vs actual", "Plots predictions against the true salaries for the best model (Gradient Boosting)."],
       ["8. Feature importance", "Plots the Random Forest's feature importance."],
       ["9. Predict a new job profile", "Encodes a custom profile and predicts a salary with the Random Forest."],
       ["10. Conclusion", "Summarises findings and suggests next steps."]], widths=[2.2, 4.5])

# ----- footer page numbers
for s_ in doc.sections:
    fp = s_.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("AI/ML Job Salary Prediction — Project Guide   |   Page "); r.font.size = Pt(9); r.font.color.rgb = SLATE
    for tag, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        run = fp.add_run(); run.font.size = Pt(9); run.font.color.rgb = SLATE
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)

doc.core_properties.title = "AI/ML Job Salary Prediction — Project Guide"
doc.core_properties.author = "Arena Agent Mode"
doc.save(OUT); print("saved", OUT, "figures:", FIG[0])
