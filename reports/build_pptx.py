"""
Builds reports/AI_ML_Salary_Prediction_Presentation.pptx
Interactive features: clickable agenda + navigation bar, native (hover-able, editable)
charts, fade transitions, click-to-reveal answers.
Run build_charts.py first (it creates the charts and metrics.json).
"""
import json, os
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "reports", "assets")
CH = os.path.join(A, "charts")
M = json.load(open(os.path.join(A, "metrics.json")))
OUT = os.path.join(ROOT, "reports", "AI_ML_Salary_Prediction_Presentation.pptx")

NAVY = "0F172A"; INDIGO = "6366F1"; TEAL = "14B8A6"; AMBER = "F59E0B"; ROSE = "F43F5E"
SLATE = "64748B"; LIGHT = "F1F5F9"; WHITE = "FFFFFF"; VIOLET = "A78BFA"; INK2 = "334155"; DARK2 = "1E293B"
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "AI/ML Job Salary Prediction"
prs.core_properties.author = "Arena Agent Mode"

def money(v): return f"${v:,.0f}"
def k(v): return f"${v/1000:,.0f}k"

# ------------------------------------------------------------------ helpers
def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = RGBColor.from_string(color)

def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: list of (text, size, bold, color) or a single tuple"""
    if isinstance(paras, tuple): paras = [paras]
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(s, paras, align, anchor)
    return s

def fill_text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.08):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for p_ in paras:
        text, size, bold, color = p_[:4]
        after = p_[4] if len(p_) > 4 else 4
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(after)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
        r.font.color.rgb = RGBColor.from_string(color)

def box(slide, x, y, w, h, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(1)
    else: s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    s.shadow.inherit = False
    return s

def card(slide, x, y, w, h, title, body, accent=INDIGO, fill=WHITE, tsize=17, bsize=13, tcolor=None):
    s = box(slide, x, y, w, h, fill=fill, line="E2E8F0")
    paras = [(title, tsize, True, tcolor or accent, 6)]
    for b in ([body] if isinstance(body, str) else body):
        paras.append((b, bsize, False, INK2, 4))
    fill_text(s, paras, margin=0.18)
    return s

def header(slide, title, subtitle=None, dark=False):
    tc = WHITE if dark else NAVY
    tb(slide, 0.6, 0.78, 12.0, 0.75, (title, 30, True, tc))
    if subtitle: tb(slide, 0.6, 1.5, 12.0, 0.45, (subtitle, 16, False, "94A3B8" if dark else SLATE))

def footer(slide, n, dark=False):
    tb(slide, 0.6, 7.0, 8, 0.3, ("AI/ML Job Salary Prediction  •  interactive deck", 9, False, "94A3B8" if dark else "94A3B8"))
    tb(slide, 11.9, 7.0, 0.9, 0.3, (str(n), 9, False, "94A3B8"), align=PP_ALIGN.RIGHT)

def link(shape, target):
    shape.click_action.target_slide = target

NAV = [("Agenda", "agenda"), ("Big Picture", "s_big"), ("Data", "s_data"), ("Explore", "s_explore"),
       ("Models", "s_models"), ("Results", "s_results"), ("Quiz", "quiz")]
def nav(slide, active, dark=False):
    x = 13.333 - 0.45 - len(NAV) * 1.14
    for label, key in NAV:
        is_active = key == active
        fill = INDIGO if is_active else ("1E293B" if dark else "E2E8F0")
        col = WHITE if (is_active or dark) else NAVY
        b = box(slide, x, 0.22, 1.08, 0.36, fill=fill, radius=0.5)
        fill_text(b, [(label, 10.5, True, col, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
        link(b, S[key])
        x += 1.14

def pic_fit(slide, path, x, y, maxw, maxh):
    w, h = Image.open(path).size
    scale = min(maxw / w, maxh / h)
    pw, ph = w * scale, h * scale
    return slide.shapes.add_picture(path, Inches(x + (maxw - pw) / 2), Inches(y + (maxh - ph) / 2), Inches(pw), Inches(ph))

def style_chart(chart, fmt=None, font=11):
    chart.font.name = FONT; chart.font.size = Pt(font); chart.font.color.rgb = RGBColor.from_string(INK2)
    chart.has_legend = False
    try:
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string("E2E8F0")
        va.format.line.fill.background()
        if fmt: va.tick_labels.number_format = fmt; va.tick_labels.number_format_is_linked = False
        chart.category_axis.format.line.color.rgb = RGBColor.from_string("CBD5E1")
    except Exception:
        pass

def bar_chart(slide, x, y, w, h, cats, vals, colors, horizontal=False, fmt='$#,##0,"k"', label_fmt=None, font=11):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series("Value", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart; style_chart(ch, fmt, font)
    plot = ch.plots[0]; plot.gap_width = 60; plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels; dl.number_format = label_fmt or fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(font); dl.font.bold = True
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = RGBColor.from_string(c)
    if horizontal:
        ch.category_axis.reverse_order = True  # first category at the top
    return ch

def pptx_timing(slide, shape_ids):
    """Click-to-reveal: each shape fades in on its own mouse click."""
    P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    n = [2]
    def nid():
        n[0] += 1; return n[0]
    pars = []
    for sid in shape_ids:
        a, b, c, d, e = nid(), nid(), nid(), nid(), nid()
        pars.append(f'''<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{c}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{e}" dur="500"/><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>''')
    bld = "".join(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>' for sid in shape_ids)
    xml = f'''<p:timing {P}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{bld}</p:bldLst></p:timing>'''
    el = etree.fromstring(xml)
    slide._element.append(el)

def transition(slide):
    P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    el = etree.fromstring(f'<p:transition {P} spd="med"><p:fade/></p:transition>')
    cm = slide._element.find(qn("p:clrMapOvr"))
    idx = list(slide._element).index(cm) + 1
    slide._element.insert(idx, el)

# ------------------------------------------------------------------ create all slides first (for links)
keys = ["title", "agenda", "s_big", "big1", "big2", "s_data", "data1", "data2", "data3", "s_explore",
        "ex1", "ex2", "ex3", "ex4", "ex5", "s_models", "mod1", "mod2", "s_results", "res1", "res2",
        "res3", "res4", "res5", "res6", "takeaways", "quiz", "end"]
S = {k_: prs.slides.add_slide(BLANK) for k_ in keys}
# Keep a stable order for numbering
ORDER = keys
num = {k_: i + 1 for i, k_ in enumerate(ORDER)}
REVEAL = []   # (slide, [shape ids]) to be animated

# ================================================================== 1 TITLE
s = S["title"]; bg(s, NAVY)
hero_path = os.path.join(A, "hero.png")
hw, hh = Image.open(hero_path).size
pic = s.shapes.add_picture(hero_path, Inches(6.9), Inches(0), Inches(6.433), Inches(7.5))
ratio_img = hw / hh; ratio_box = 6.433 / 7.5
if ratio_img > ratio_box:
    cut = (1 - ratio_box / ratio_img) / 2; pic.crop_left = cut; pic.crop_right = cut
tb(s, 0.7, 1.55, 6, 0.4, ("DATA SCIENCE  •  MACHINE LEARNING  •  BEGINNER FRIENDLY", 13, True, TEAL))
tb(s, 0.7, 2.0, 6.1, 2.0, ("AI/ML Job Salary Prediction", 46, True, WHITE))
tb(s, 0.7, 3.65, 6.0, 1.2, ("How 500 AI & machine-learning job records are used to teach a computer to estimate salaries, explained step by step.", 18, False, "CBD5E1"))
b1 = box(s, 0.7, 5.0, 2.9, 0.6, fill=TEAL, radius=0.5)
fill_text(b1, [("Start the tour  →", 15, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
link(b1, S["agenda"])
b2 = box(s, 3.8, 5.0, 2.6, 0.6, fill="1E293B", radius=0.5)
fill_text(b2, [("Jump to quiz", 15, True, WHITE, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
link(b2, S["quiz"])
tb(s, 0.7, 6.6, 6, 0.4, ("Based on the GitHub project by Vismay Vinod  •  MIT License", 11, False, "94A3B8"))
transition(s)

# ================================================================== 2 AGENDA
s = S["agenda"]; bg(s, LIGHT); nav(s, "agenda")
header(s, "Today's journey", "Click any card to jump straight to that part of the deck")
items = [("01", "The Big Picture", "What the project does, in plain words", "s_big", INDIGO),
         ("02", "The Data", "The 500 job records and how clean they are", "s_data", TEAL),
         ("03", "Exploring Patterns", "What affects salary? Charts & numbers", "s_explore", VIOLET),
         ("04", "Building Models", "The four 'learning' models and how we prepared data", "s_models", AMBER),
         ("05", "Results & Takeaways", "Which model won, how accurate it is, and why", "s_results", ROSE),
         ("✓", "Test yourself", "Three quick questions with click-to-reveal answers", "quiz", INDIGO)]
for i, (n_, t_, d_, tgt, col) in enumerate(items):
    cx = 0.6 + (i % 3) * 4.1; cy = 2.25 + (i // 3) * 2.35
    c = box(s, cx, cy, 3.85, 2.05, fill=WHITE, line="E2E8F0", radius=0.08)
    fill_text(c, [(n_, 30, True, col, 2), (t_, 18, True, NAVY, 4), (d_, 13, False, SLATE, 0)], margin=0.25)
    link(c, S[tgt])
footer(s, num["agenda"])

# ================================================================== SECTION DIVIDERS
def divider(slide, n_, title, desc, color, key):
    bg(slide, NAVY); nav(slide, key, dark=True)
    tb(slide, 0.9, 2.2, 6, 1.6, (n_, 96, True, color))
    tb(slide, 0.9, 3.9, 11, 0.9, (title, 40, True, WHITE))
    tb(slide, 0.9, 4.85, 10.5, 0.8, (desc, 18, False, "94A3B8"))
    b = box(slide, 0.9, 5.9, 2.6, 0.5, fill=DARK2, radius=0.5)
    fill_text(b, [("← Back to agenda", 12, True, WHITE, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    link(b, S["agenda"])
    transition(slide)
divider(S["s_big"], "01", "The Big Picture", "What this project is about, and how a computer learns from examples.", INDIGO, "s_big")
divider(S["s_data"], "02", "The Data", "Where the numbers come from, what each column means, and whether we can trust them.", TEAL, "s_data")
divider(S["s_explore"], "03", "Exploring Patterns", "Before modelling, look at the data: what raises or lowers salary?", VIOLET, "s_explore")
divider(S["s_models"], "04", "Building the Models", "Preparing the data, splitting it, and teaching four different models.", AMBER, "s_models")
divider(S["s_results"], "05", "Results & Takeaways", "Which model predicts best, how big the errors are, and what we learned.", ROSE, "s_results")

# ================================================================== 01 BIG PICTURE
s = S["big1"]; bg(s, LIGHT); nav(s, "s_big"); header(s, "What is this project?", "A simple way to say: 'here is what a similar AI/ML job usually pays'")
card(s, 0.6, 2.2, 3.9, 2.6, "The problem", ["Job-seekers and companies want a fair idea of what an AI/ML role should pay.", "Salary depends on many things at once."], INDIGO)
card(s, 4.72, 2.2, 3.9, 2.6, "The data", ["500 job records.", "Each record has 11 facts: experience, education, skills, city, company and more."], TEAL)
card(s, 8.84, 2.2, 3.9, 2.6, "The method", ["Train computer models on past jobs.", "Then ask them to predict salaries for jobs they have never seen."], AMBER)
c = box(s, 0.6, 5.1, 12.14, 1.6, fill=NAVY, radius=0.08)
fill_text(c, [("Everyday analogy", 14, True, TEAL, 6),
              ("A student studies 400 solved salary problems, then is tested on 100 new ones they have never seen. "
               "If they score well, we trust them to answer the next question. That is exactly what this project does.", 17, False, WHITE, 0)],
          margin=0.3, anchor=MSO_ANCHOR.MIDDLE)
footer(s, num["big1"]); transition(s)

s = S["big2"]; bg(s, WHITE); nav(s, "s_big"); header(s, "The seven steps, top to bottom", "Every part of the notebook follows this same path")
pic_fit(s, os.path.join(CH, "c16_workflow.png"), 0.6, 2.2, 12.1, 3.4)
tb(s, 0.6, 5.9, 12.1, 0.9, [("Think of it like cooking: gather ingredients (data), taste them (explore), wash and chop (prepare), "
                             "keep some for a test dish (split), cook several recipes (train), taste-test (evaluate), then serve (predict).", 15, False, INK2, 0)])
footer(s, num["big2"]); transition(s)

# ================================================================== 02 DATA
s = S["data1"]; bg(s, LIGHT); nav(s, "s_data"); header(s, "The dataset at a glance", "ai_ml_job_analysis.csv: one row per job, one column per fact")
kpis = [(f"{M['rows']}", "job records", INDIGO), (f"{M['cols']}", "columns (11 facts each)", TEAL),
        (f"{M['n_titles']}", "different job titles", VIOLET), (f"{M['n_companies']}", "employers", AMBER),
        (money(M["salary_mean"]), "average salary (USD)", ROSE), (f"{k(M['salary_min'])}–{k(M['salary_max'])}", "salary range", INDIGO)]
for i, (big, lab, col) in enumerate(kpis):
    x = 0.6 + (i % 3) * 4.1; y = 2.2 + (i // 3) * 2.3
    c = box(s, x, y, 3.85, 2.0, fill=WHITE, line="E2E8F0", radius=0.08)
    box(s, x, y, 0.1, 2.0, fill=col, radius=0.0)
    fill_text(c, [(big, 36, True, col, 2), (lab, 15, False, SLATE, 0)], margin=0.3, anchor=MSO_ANCHOR.MIDDLE)
footer(s, num["data1"]); transition(s)

s = S["data2"]; bg(s, WHITE); nav(s, "s_data"); header(s, "Meet the 11 columns", "What each column means, with a real example from row 3 of the file")
rows = [("Column", "What it means", "Example (row 3)"),
        ("ID", "Row number (ignored by the model)", "3"),
        ("Job Title", "Role name", "AI Consultant"),
        ("Experience (Years)", "Years of work experience", "13.0"),
        ("Education", "Highest degree: Bachelors / Masters / PhD", "Masters"),
        ("Skill Count", "How many technical skills are listed", "16"),
        ("Industry", "Sector the job belongs to", "Technology"),
        ("Company", "Employer name", "Cognizant"),
        ("Location", "City or country (15 values)", "London, UK"),
        ("Remote Work", "Yes / No / Hybrid", "No"),
        ("Certification", "Professional certificate held, or 'None'", "IBM Data Science Cert"),
        ("Salary (USD)", "THE TARGET: the number we want to predict", "143,200")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(2.1), Inches(12.1), Inches(4.7)).table
tbl.columns[0].width = Inches(2.8); tbl.columns[1].width = Inches(5.6); tbl.columns[2].width = Inches(3.7)
for r, row in enumerate(rows):
    for c_, val in enumerate(row):
        cell = tbl.cell(r, c_); cell.text = ""
        cell.fill.solid()
        if r == 0: cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
        elif row[0] == "Salary (USD)": cell.fill.fore_color.rgb = RGBColor.from_string("FEF3C7")
        else: cell.fill.fore_color.rgb = RGBColor.from_string(WHITE if r % 2 else LIGHT)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        p = cell.text_frame.paragraphs[0]; run = p.add_run(); run.text = val
        run.font.size = Pt(12.5); run.font.name = FONT
        run.font.bold = (r == 0 or c_ == 0 or row[0] == "Salary (USD)")
        run.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else NAVY if c_ == 0 else INK2)
footer(s, num["data2"]); transition(s)

s = S["data3"]; bg(s, LIGHT); nav(s, "s_data"); header(s, "Where does the salary sit?", "Most jobs pay between $41k and $134k. Few jobs pay very high salaries.")
hist_bins = list(range(0, 240001, 20000))
import numpy as np
import pandas as pd
sal = pd.read_csv(os.path.join(ROOT, "ai_ml_job_analysis.csv"))["Salary (USD)"]
counts, _ = np.histogram(sal, bins=hist_bins)
cats = [f"{k(a)}–{k(b)}" for a, b in zip(hist_bins[:-1], hist_bins[1:])]
ch = bar_chart(s, 0.6, 2.1, 8.2, 4.7, cats, [int(c_) for c_ in counts], [INDIGO] * len(counts), fmt="0", label_fmt="0", font=10)
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "Number of jobs"
ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
card(s, 9.2, 2.2, 3.55, 1.45, "Average vs typical", [f"Average {money(M['salary_mean'])}", f"Median {money(M['salary_median'])}"], INDIGO, bsize=14)
card(s, 9.2, 3.85, 3.55, 1.45, "Why they differ", "A few very high salaries pull the average upward. This is called a 'right skew'.", TEAL, bsize=13)
card(s, 9.2, 5.5, 3.55, 1.3, "Hover me!", "Click the chart to see the exact count for each bin.", AMBER, bsize=13)
footer(s, num["data3"]); transition(s)

# ================================================================== 03 EXPLORE
s = S["ex1"]; bg(s, WHITE); nav(s, "s_explore")
header(s, "Location is the biggest driver", "Average salary by city/country (hover a bar for the exact value)")
locs = list(M["loc_means"].keys()); lv = [M["loc_means"][l] for l in locs]
REG = {"India": ROSE, "USA": INDIGO, "Remote": AMBER}
def rcol(l):
    if "India" in l: return REG["India"]
    if "USA" in l: return REG["USA"]
    if l == "Remote": return REG["Remote"]
    return TEAL
bar_chart(s, 0.5, 2.0, 7.6, 4.9, locs, lv, [rcol(l) for l in locs], horizontal=True, font=10)
rm = M["region_means"]; rc = M["region_counts"]
for i, (rg, col) in enumerate([("USA", INDIGO), ("Europe / Canada / Singapore", TEAL), ("Remote", AMBER), ("India", ROSE)]):
    card(s, 8.5, 2.1 + i * 1.22, 4.25, 1.1, f"{rg}: {k(rm[rg])} average",
         f"{rc[rg]} jobs in this group", col, tsize=15, bsize=12)
footer(s, num["ex1"]); transition(s)

s = S["ex2"]; bg(s, LIGHT); nav(s, "s_explore")
header(s, "More experience, higher salary", "Each dot is one job. The pink line shows the overall direction.")
cd = XyChartData()
for e_, col in [("Bachelors", TEAL), ("Masters", INDIGO), ("PhD", VIOLET)]:
    pass
df_ = pd.read_csv(os.path.join(ROOT, "ai_ml_job_analysis.csv"))
for e_ in ["Bachelors", "Masters", "PhD"]:
    ser = cd.add_series(e_)
    for _, r_ in df_[df_.Education == e_].iterrows(): ser.add_data_point(r_["Experience (Years)"], r_["Salary (USD)"])
trend = cd.add_series("Overall trend")
z = np.polyfit(df_["Experience (Years)"], df_["Salary (USD)"], 1)
for xv in [0, 20]: trend.add_data_point(xv, float(np.polyval(z, xv)))
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(0.5), Inches(2.0), Inches(7.9), Inches(4.9), cd)
ch = gf.chart; style_chart(ch, '$#,##0,"k"', 11); ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False; ch.legend.font.size = Pt(11)
ch.category_axis.has_title = True; ch.category_axis.axis_title.text_frame.text = "Years of experience"
ch.category_axis.tick_labels.number_format = "0"; ch.category_axis.tick_labels.number_format_is_linked = False
for ser, col in zip(ch.series, [TEAL, INDIGO, VIOLET, ROSE]):
    if ser.name == "Overall trend":
        ser.marker.style = XL_MARKER_STYLE.NONE; ser.format.line.color.rgb = RGBColor.from_string(ROSE)
        ser.format.line.width = Pt(2.75); ser.smooth = False
    else:
        ser.format.line.fill.background(); ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 8
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(col)
        ser.marker.format.line.color.rgb = RGBColor.from_string(WHITE)
ch.has_legend = True
bands = M["exp_bands"]
tb(s, 8.7, 2.05, 4.1, 0.4, ("Average salary by experience", 15, True, NAVY))
for i, (bn, bv) in enumerate(bands.items()):
    card(s, 8.7, 2.5 + i * 0.86, 4.05, 0.74, f"{bn}:  {k(bv['mean'])}", f"{bv['count']} jobs", [INDIGO, TEAL, VIOLET, AMBER, ROSE][i], tsize=14, bsize=11)
tb(s, 8.7, 6.8, 4.1, 0.3, (f"Correlation r = {M['corr_exp_salary']:.2f} (a moderate positive link)", 11, False, SLATE))
footer(s, num["ex2"]); transition(s)

s = S["ex3"]; bg(s, WHITE); nav(s, "s_explore")
header(s, "Education, remote work and certificates", "Two quick charts, then a short note on certificates")
em = M["edu_means"]; edus = ["Bachelors", "Masters", "PhD"]
tb(s, 0.6, 2.05, 6, 0.4, ("Average salary by education", 15, True, NAVY))
bar_chart(s, 0.5, 2.45, 6.0, 3.9, edus, [em[e] for e in edus], [TEAL, INDIGO, VIOLET], font=12)
rmn = M["remote_means"]; rcs = ["No", "Hybrid", "Yes"]
tb(s, 6.9, 2.05, 6, 0.4, ("Average salary by remote work", 15, True, NAVY))
bar_chart(s, 6.8, 2.45, 6.0, 3.9, ["Not remote", "Hybrid", "Remote"], [rmn[c_] for c_ in rcs], [SLATE, INDIGO, TEAL], font=12)
cm_ = M["cert_means"]; best_c = max(cm_, key=cm_.get); worst_c = min(cm_, key=cm_.get)
tb(s, 0.6, 6.4, 12.1, 0.55, [(f"Takeaway: a PhD earns about {k(em['PhD'])} on average vs {k(em['Bachelors'])} for a Bachelors. "
                                f"Remote work barely changes pay. Certificates vary from {k(cm_[worst_c])} ({worst_c}) to {k(cm_[best_c])} ({best_c}).", 13, False, INK2, 0)])
footer(s, num["ex3"]); transition(s)

s = S["ex4"]; bg(s, LIGHT); nav(s, "s_explore")
header(s, "Job titles: role matters", "Average salary per job title, sorted from highest to lowest")
tm = M["title_means"]; tt = list(tm.keys()); tv = list(tm.values())
bar_chart(s, 0.5, 2.0, 8.3, 5.0, tt, tv, [INDIGO if v >= 80000 else ROSE for v in tv], horizontal=True, font=10)
card(s, 9.1, 2.2, 3.65, 2.1, f"Top: {tt[0]}", f"Average {money(tv[0])}. Often senior, strategy-focused roles.", INDIGO, bsize=13)
card(s, 9.1, 4.45, 3.65, 2.1, f"Lowest: {tt[-1]}", f"Average {money(tv[-1])}. Typically early-career roles in this data.", ROSE, bsize=13)
footer(s, num["ex4"]); transition(s)

s = S["ex5"]; bg(s, WHITE); nav(s, "s_explore")
header(s, "How numbers move together", "A correlation score runs from -1 to +1. Closer to 0 means little link.")
pic_fit(s, os.path.join(CH, "c10_correlation.png"), 0.6, 2.05, 5.8, 4.8)
cards_txt = [("Experience vs salary: +0.36", "A moderate positive link. More experience usually means higher pay."),
             ("Skill count vs salary: +0.13", "A weak link. Listing more skills does not guarantee a big salary."),
             ("Experience vs skill count: -0.02", "Basically unrelated: the two carry different information for the model.")]
for i, (t_, b_) in enumerate(cards_txt):
    card(s, 6.8, 2.1 + i * 1.12, 5.95, 0.98, t_, b_, [INDIGO, TEAL, AMBER][i], tsize=14, bsize=12)
q = box(s, 6.8, 5.5, 5.95, 0.6, fill=DARK2, radius=0.08)
fill_text(q, [("Quick question: what does -0.02 mean?  (click the box below)", 13, True, WHITE, 0)], margin=0.2, anchor=MSO_ANCHOR.MIDDLE)
ans = box(s, 6.8, 6.2, 5.95, 0.65, fill=TEAL, radius=0.08)
fill_text(ans, [("Answer: almost no straight-line relationship. The two numbers do not move together.", 12.5, False, NAVY, 0)], margin=0.2, anchor=MSO_ANCHOR.MIDDLE)
REVEAL.append((s, [ans.shape_id]))
footer(s, num["ex5"]); transition(s)

# ================================================================== 04 MODELS
s = S["mod1"]; bg(s, LIGHT); nav(s, "s_models")
header(s, "Preparing the data", "Computers only understand numbers, so words must be converted first")
rows = [("Step", "What we do", "Why"),
        ("1. Drop ID", "Remove the row number", "It has no link to salary"),
        ("2. Education", "Bachelors = 0, Masters = 1, PhD = 2", "Keeps the natural order of degrees"),
        ("3. Remote Work", "Turn Yes / No / Hybrid into three 0-or-1 columns", "These options have no natural order"),
        ("4. Text columns", "Job Title, Industry, Company, Location, Certification get a number each (LabelEncoder)", "Models need numbers. Note: this order is arbitrary (see limitations)"),
        ("5. Scale numbers", "Put numbers on a similar range (linear model only)", "Helps the straight-line model work well")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(2.1), Inches(12.1), Inches(4.4)).table
tbl.columns[0].width = Inches(2.6); tbl.columns[1].width = Inches(5.4); tbl.columns[2].width = Inches(4.1)
for r, row in enumerate(rows):
    for c_, val in enumerate(row):
        cell = tbl.cell(r, c_); cell.text = ""; cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else (WHITE if r % 2 else LIGHT))
        p = cell.text_frame.paragraphs[0]; run = p.add_run(); run.text = val
        run.font.size = Pt(13); run.font.name = FONT; run.font.bold = (r == 0 or c_ == 0)
        run.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else NAVY if c_ == 0 else INK2)
footer(s, num["mod1"]); transition(s)

s = S["mod2"]; bg(s, WHITE); nav(s, "s_models")
header(s, "Four models, four ways to learn", "Same data, same split. Each model learns the pattern differently.")
models = [("Linear Regression", "Draws one straight line through the data. Fast and simple. The baseline everyone compares against.", "Analogy: a ruler", INDIGO),
          ("Decision Tree", "Asks a chain of yes/no questions such as 'Is experience above 10 years?'", "Analogy: a flowchart", TEAL),
          ("Random Forest", "300 trees are each trained on a random part of the data, then their answers are averaged.", "Analogy: asking 300 experts and averaging", VIOLET),
          ("Gradient Boosting", "Trees are built one after another; each new tree fixes the mistakes of the previous ones.", "Analogy: correcting your own homework again and again", AMBER)]
for i, (t_, d_, an, col) in enumerate(models):
    x = 0.6 + i * 3.05
    c = box(s, x, 2.15, 2.85, 3.85, fill=WHITE, line="E2E8F0", radius=0.07)
    box(s, x, 2.15, 2.85, 0.12, fill=col, radius=0.0)
    fill_text(c, [(t_, 17, True, col, 8), (d_, 13, False, INK2, 10), (an, 12, True, SLATE, 0)], margin=0.2)
c = box(s, 0.6, 6.2, 12.1, 0.62, fill=NAVY, radius=0.5)
fill_text(c, [("Train/test split:  400 jobs to learn from  •  100 jobs hidden for testing  (80/20, random_state = 42)", 14, True, WHITE, 0)],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, num["mod2"]); transition(s)

# ================================================================== 05 RESULTS
s = S["res1"]; bg(s, LIGHT); nav(s, "s_results")
R = M["results"]; names = ["Linear Regression", "Decision Tree", "Random Forest", "Gradient Boosting"]
header(s, "Which model predicts best?", "Higher R² and lower errors are better. Gradient Boosting wins on every measure.")
bar_chart(s, 0.5, 2.05, 6.2, 4.85, ["Linear", "Decision Tree", "Random Forest", "Gradient Boosting"],
          [R[n]["R2"] for n in names], [SLATE, AMBER, INDIGO, TEAL], fmt="0.00", label_fmt="0.000", font=12)
tbl = s.shapes.add_table(5, 4, Inches(7.0), Inches(2.15), Inches(5.75), Inches(2.9)).table
hdr = ["Model", "R²", "MAE", "RMSE"]
for c_, h_ in enumerate(hdr):
    cell = tbl.cell(0, c_); cell.text = ""; cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    rr = cell.text_frame.paragraphs[0].add_run(); rr.text = h_; rr.font.bold = True; rr.font.size = Pt(13); rr.font.color.rgb = RGBColor.from_string(WHITE); rr.font.name = FONT
for r, n in enumerate(names, start=1):
    vals = [n, f"{R[n]['R2']:.3f}", money(R[n]["MAE"]), money(R[n]["RMSE"])]
    for c_, v in enumerate(vals):
        cell = tbl.cell(r, c_); cell.text = ""; cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string("CCFBF1" if n == M["best_model"] else WHITE)
        rr = cell.text_frame.paragraphs[0].add_run(); rr.text = v; rr.font.size = Pt(13); rr.font.name = FONT
        rr.font.bold = n == M["best_model"]; rr.font.color.rgb = RGBColor.from_string(NAVY)
tbl.columns[0].width = Inches(2.35); [setattr(tbl.columns[i], "width", Inches(1.1)) for i in (1, 2, 3)]
card(s, 7.0, 5.35, 5.75, 1.45, "What do these mean?",
     [f"R² = how much of the salary variation the model explains (1.0 = perfect).",
      f"MAE = on average, the guess is off by about {money(R[M['best_model']]['MAE'])}."], TEAL, tsize=14, bsize=12)
footer(s, num["res1"]); transition(s)

s = S["res2"]; bg(s, WHITE); nav(s, "s_results")
header(s, "How close are the predictions?", "Left: each dot is a test job. Dots near the dashed line were predicted correctly.")
pic_fit(s, os.path.join(CH, "c13_pred_vs_actual.png"), 0.4, 2.05, 6.4, 4.85)
pic_fit(s, os.path.join(CH, "c14_residuals.png"), 6.9, 2.05, 6.0, 3.5)
tb(s, 6.9, 5.6, 6.0, 1.3, [(f"{M['resid_within_20k']*100:.0f}% of the test predictions were within $20,000 of the real salary.", 16, True, NAVY, 6),
                           ("Errors (right) are centred near zero, so the model is not systematically too high or too low.", 13, False, INK2, 0)])
footer(s, num["res2"]); transition(s)

s = S["res3"]; bg(s, LIGHT); nav(s, "s_results")
header(s, "Is the result lucky? Cross-validation says no", "We repeated the test 5 times on 5 different splits of the data")
cvv = M["cv"]
bar_chart(s, 0.5, 2.05, 7.2, 4.8, ["Linear", "Decision Tree", "Random Forest", "Gradient Boosting"],
          [cvv[n]["mean"] for n in names], [SLATE, AMBER, INDIGO, TEAL], fmt="0.00", label_fmt="0.00", font=12)
card(s, 8.0, 2.2, 4.75, 1.5, "How to read it", "Cross-validation shares the data into 5 parts and rotates which part is the test set.", INDIGO, bsize=13)
card(s, 8.0, 3.9, 4.75, 1.5, "What we found", f"Gradient Boosting averages R² {cvv['Gradient Boosting']['mean']:.2f} (± {cvv['Gradient Boosting']['std']:.3f}), the highest and steadiest.", TEAL, bsize=13)
card(s, 8.0, 5.6, 4.75, 1.25, "Simple summary", "The winner is not a fluke of one lucky split.", AMBER, bsize=13)
footer(s, num["res3"]); transition(s)

s = S["res4"]; bg(s, WHITE); nav(s, "s_results")
header(s, "What does the model rely on?", "Feature importance from the Random Forest (share of its decisions)")
imp = M["importance_rf"]; fk = list(imp.keys()); fv = list(imp.values())
def fcol(f):
    if f.startswith("Remote"): return AMBER
    if f in ("Experience (Years)", "Skill Count", "Education"): return TEAL
    return INDIGO
bar_chart(s, 0.5, 2.0, 7.6, 5.0, [f.replace("_", " ") for f in fk], fv, [fcol(f) for f in fk], horizontal=True, fmt="0%", label_fmt="0.0%", font=10)
card(s, 8.4, 2.2, 4.35, 1.5, f"Location: {imp['Location']*100:.0f}%", "The single most important input. Where the job is matters most.", INDIGO, tsize=16, bsize=12)
card(s, 8.4, 3.9, 4.35, 1.5, f"Experience: {imp['Experience (Years)']*100:.0f}%", "The strongest number-based factor, ahead of skill count.", TEAL, tsize=16, bsize=12)
card(s, 8.4, 5.6, 4.35, 1.25, "Remote work: under 1% each", "Surprisingly little influence on salary in this data.", AMBER, tsize=15, bsize=12)
footer(s, num["res4"]); transition(s)

s = S["res5"]; bg(s, LIGHT); nav(s, "s_results")
header(s, "Try a prediction", "A new candidate profile, predicted by two models")
prof = M["example_profile"]
c = box(s, 0.6, 2.1, 4.9, 4.7, fill=WHITE, line="E2E8F0", radius=0.06)
lines = [("Example candidate", 16, True, INDIGO, 8)]
for kk, vv in prof.items(): lines.append((f"{kk}:  {vv}", 13, False, INK2, 3))
fill_text(c, lines, margin=0.3)
c2 = box(s, 5.8, 2.1, 3.4, 2.25, fill=INDIGO, radius=0.07)
fill_text(c2, [("Random Forest", 14, True, WHITE, 4), (money(M["example_pred_rf"]), 30, True, WHITE, 2), ("with 'Bengaluru' (see note)", 11, False, "E0E7FF", 0)], margin=0.25, anchor=MSO_ANCHOR.MIDDLE)
c3 = box(s, 9.4, 2.1, 3.35, 2.25, fill=TEAL, radius=0.07)
fill_text(c3, [("Gradient Boosting", 14, True, NAVY, 4), (money(M["example_pred_gb"]), 30, True, NAVY, 2), ("with 'Bengaluru' (see note)", 11, False, NAVY, 0)], margin=0.25, anchor=MSO_ANCHOR.MIDDLE)
c4 = box(s, 5.8, 4.55, 6.95, 2.25, fill=NAVY, radius=0.07)
fill_text(c4, [("Important catch", 15, True, AMBER, 6),
               (f"'Bengaluru, India' is not in the dataset (it is spelled 'Bangalore, India'). The notebook silently treats it as '{M['example_fallback_location']}'.", 13, False, WHITE, 6),
               (f"With the correct spelling: Random Forest {money(M['example_pred_rf_fixed'])}, Gradient Boosting {money(M['example_pred_gb_fixed'])}.", 14, True, WHITE, 0)],
          margin=0.3, anchor=MSO_ANCHOR.MIDDLE)
footer(s, num["res5"]); transition(s)

s = S["res6"]; bg(s, WHITE); nav(s, "s_results")
header(s, "Limits: what to keep in mind", "A good score is not the same as a perfect tool")
lim = [("Small dataset", "500 rows. Results could change with more data.", ROSE),
       ("Data source unclear", "Some labels look odd (e.g. many 'Research Interns' with 5+ years). It may be synthetic.", AMBER),
       ("Arbitrary codes", "LabelEncoder gives text values numbers in alphabetical order, which the model may misread.", VIOLET),
       ("Unseen names", "A city missing from the data silently falls back to the first known one (see the Bengaluru example).", ROSE),
       ("Correlation ≠ cause", "Patterns show association. They do not prove that a skill or city causes a higher salary.", TEAL),
       ("Single split", f"The notebook evaluates one 80/20 split. Cross-validation (slide {num['res3']}) adds more confidence.", INDIGO)]
for i, (t_, d_, col) in enumerate(lim):
    x = 0.6 + (i % 3) * 4.1; y = 2.1 + (i // 3) * 2.4
    card(s, x, y, 3.85, 2.15, t_, d_, col, tsize=16, bsize=13)
footer(s, num["res6"]); transition(s)

s = S["takeaways"]; bg(s, NAVY); nav(s, "s_results", dark=True)
tb(s, 0.6, 0.78, 12, 0.75, ("Key takeaways", 30, True, WHITE))
tk = [("1", "Location matters most", "US & Europe-based roles average far above India-based roles in this data.", INDIGO),
      ("2", "Experience pays", "Average salary climbs from about $65k (0-3 yrs) to $120k (15-20 yrs).", TEAL),
      ("3", "Boosting wins", f"Gradient Boosting: R² {R['Gradient Boosting']['R2']:.3f}, average error about {money(R['Gradient Boosting']['MAE'])}.", AMBER),
      ("4", "Check the data", "Clean labels and clear city names matter more than fancy models.", ROSE),
      ("5", "Start simple", "Begin with a baseline, then move to better models. Compare them honestly.", VIOLET)]
for i, (n_, t_, d_, col) in enumerate(tk):
    y = 1.85 + i * 1.0
    b = box(s, 0.6, y, 0.8, 0.8, fill=col, shape=MSO_SHAPE.OVAL)
    fill_text(b, [(n_, 22, True, WHITE, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    tb(s, 1.65, y - 0.02, 11, 0.45, (t_, 19, True, WHITE))
    tb(s, 1.65, y + 0.36, 11, 0.45, (d_, 13.5, False, "CBD5E1"))
footer(s, num["takeaways"], dark=True); transition(s)

# ================================================================== QUIZ
s = S["quiz"]; bg(s, LIGHT); nav(s, "quiz")
header(s, "Test yourself", "Think of the answer first, then click to reveal it")
qs = [("Q1. If a model's R² is 0.93 on unseen jobs, is that good?",
       "Yes. R² is close to 1, so the model explains about 93% of the salary variation on jobs it never saw."),
      ("Q2. Which input mattered most to the Random Forest?",
       f"Location, with about {M['importance_rf']['Location']*100:.0f}% of its decisions. Experience came second."),
      ("Q3. Why do we keep 100 jobs hidden during training?",
       "To check the model on jobs it has never seen, so we do not fool ourselves with memorised answers.")]
for i, (qq, aa) in enumerate(qs):
    y = 2.1 + i * 1.6
    qb = box(s, 0.6, y, 5.6, 1.35, fill=WHITE, line="E2E8F0", radius=0.08)
    fill_text(qb, [(qq, 16, True, NAVY, 0)], margin=0.25, anchor=MSO_ANCHOR.MIDDLE)
    ab = box(s, 6.5, y, 6.25, 1.35, fill=TEAL, radius=0.08)
    fill_text(ab, [("Answer: " + aa, 14, False, NAVY, 0)], margin=0.25, anchor=MSO_ANCHOR.MIDDLE)
    REVEAL.append((s, [ab.shape_id]))
footer(s, num["quiz"]); transition(s)

# ================================================================== END
s = S["end"]; bg(s, NAVY); nav(s, "agenda", dark=True)
tb(s, 0.9, 2.0, 11, 1.0, ("Thank you!", 50, True, WHITE))
tb(s, 0.9, 3.1, 11, 0.8, ("Next steps to explore: GridSearchCV tuning, XGBoost/LightGBM, and a Streamlit app for live predictions.", 18, False, "CBD5E1"))
tb(s, 0.9, 4.05, 11, 0.5, ("Source: github.com/Vismay-dev1/AI-ML-Salary-Prediction", 14, False, TEAL))
b = box(s, 0.9, 5.0, 2.8, 0.6, fill=TEAL, radius=0.5)
fill_text(b, [("Back to agenda", 14, True, NAVY, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE); link(b, S["agenda"])
b = box(s, 4.0, 5.0, 2.6, 0.6, fill="1E293B", radius=0.5)
fill_text(b, [("Retake quiz", 14, True, WHITE, 0)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE); link(b, S["quiz"])
transition(s)

# ------------------------------------------------------------------ apply click-to-reveal animations
seen = {}
for sl, ids in REVEAL:
    seen.setdefault(id(sl), (sl, []))[1].extend(ids)
for sl, ids in seen.values():
    pptx_timing(sl, ids)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
