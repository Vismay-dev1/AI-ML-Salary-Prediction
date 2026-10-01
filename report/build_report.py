# -*- coding: utf-8 -*-
"""Builds the DOCX and the PDF of the report from ONE content model (report_content.B).
   PDF  : reportlab (Times-compatible STIX font, embedded)  -> real page numbers for the TOC
   DOCX : python-docx (Times New Roman / Courier New)       -> TOC field pre-filled with the PDF page numbers
Run:  python build_report.py"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content import B
from content_base import HERE, REPO, A
from PIL import Image as PILImage

OUTDIR = HERE
BASENAME = "AI_ML_Job_Salary_Prediction_Internship_Demonstration"
TITLE = "AI/ML Job Salary Prediction"
HEADER_TXT = "AI/ML Job Salary Prediction — Internship Project Demonstration"
FOOTER_TXT = "Vismay Vinod  |  BCA – Artificial Intelligence and Machine Learning Internship"
REPO_URL = "https://github.com/Vismay-dev1/AI-ML-Salary-Prediction"

COVER_ROWS = [("Submitted By:", "Vismay Vinod"), ("Course:", "Bachelor of Computer Applications (BCA)"),
              ("Semester:", "5th / 6th Semester (as applicable)"), ("Internship Area:", "Artificial Intelligence and Machine Learning"),
              ("Project Type:", "Machine Learning Regression Project"), ("Academic Year:", "2026–2027"),
              ("College Name:", "[College Name]"), ("University Name:", "[University Name]"), ("Department:", "[Department Name]"),
              ("Internship Organization:", "[Internship Organization]"), ("Faculty / Guide Name:", "[Guide Name]"),
              ("Date of Submission:", "[DD / MM / YYYY]")]

TOKEN = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`")
def tokens(text):
    """-> list of (text, bold, italic, code)"""
    out, pos = [], 0
    for mt in TOKEN.finditer(text):
        if mt.start() > pos: out.append((text[pos:mt.start()], False, False, False))
        if mt.group(1) is not None: out.append((mt.group(1), True, False, False))
        elif mt.group(2) is not None: out.append((mt.group(2), False, True, False))
        else: out.append((mt.group(3), False, False, True))
        pos = mt.end()
    if pos < len(text): out.append((text[pos:], False, False, False))
    return out

def plain(text): return "".join(t for t, *_ in tokens(text))

def is_long(text):
    return any(len(w) > 24 for w in plain(text).split())

def img_size_cm(path, width_cm):
    w, h = PILImage.open(path).size
    if width_cm is None:
        dpi = 200.0
        width_cm = min(w / dpi * 2.54, 16.4)
    hc = width_cm * h / w
    if hc > 19.0:
        width_cm *= 19.0 / hc; hc = 19.0
    return width_cm, hc

# =============================================================================================
#                                           PDF
# =============================================================================================
def build_pdf(path):
    import matplotlib
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, PageBreak, Table, TableStyle,
                                    Image, KeepTogether, Flowable, CondPageBreak)
    from reportlab.platypus.tableofcontents import TableOfContents

    fdir = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
    for name, fn in [("Body", "STIXGeneral.ttf"), ("Body-Bold", "STIXGeneralBol.ttf"), ("Body-Italic", "STIXGeneralItalic.ttf"),
                     ("Body-BoldItalic", "STIXGeneralBolIta.ttf"), ("Code", "DejaVuSansMono.ttf"), ("Code-Bold", "DejaVuSansMono-Bold.ttf")]:
        pdfmetrics.registerFont(TTFont(name, os.path.join(fdir, fn)))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold", italic="Body-Italic", boldItalic="Body-BoldItalic")
    pdfmetrics.registerFontFamily("Code", normal="Code", bold="Code-Bold", italic="Code", boldItalic="Code-Bold")

    W, H = A4
    LM = RM = 2.3 * cm; TM = 2.7 * cm; BM = 2.5 * cm
    FW = W - LM - RM
    ink = colors.HexColor("#111111")

    def S(name, **kw):
        base = dict(fontName="Body", fontSize=14, leading=18.5, textColor=ink)
        base.update(kw); return ParagraphStyle(name, **base)
    st_body = S("body", alignment=TA_JUSTIFY, spaceAfter=7)
    st_body_l = S("bodyl", alignment=TA_LEFT, spaceAfter=7)
    st_h = dict(fontName="Body-Bold", fontSize=18, leading=23)
    st_h1 = S("H1", **st_h, spaceBefore=0, spaceAfter=12, keepWithNext=1)
    st_h1c = S("H1c", **st_h, spaceBefore=0, spaceAfter=14, keepWithNext=1, alignment=TA_CENTER)
    st_h2 = S("H2", **st_h, spaceBefore=14, spaceAfter=8, keepWithNext=1)
    st_h3 = S("H3", **st_h, spaceBefore=10, spaceAfter=6, keepWithNext=1)
    st_bul = S("bul", alignment=TA_JUSTIFY, leftIndent=20, bulletIndent=6, spaceAfter=4)
    st_bul_l = S("bull", alignment=TA_LEFT, leftIndent=20, bulletIndent=6, spaceAfter=4)
    st_note = S("note", fontName="Body-Italic", fontSize=12, leading=15.5, alignment=TA_JUSTIFY, backColor=colors.HexColor("#F2F2F2"),
                borderPadding=(5, 6, 5, 6), spaceBefore=6, spaceAfter=12, leftIndent=6, rightIndent=6)
    st_cap = S("cap", fontSize=11.5, leading=14.5, alignment=TA_CENTER, spaceBefore=4, spaceAfter=12)
    st_qa_q = S("qaq", fontName="Body-Bold", keepWithNext=1, spaceAfter=2, spaceBefore=4)
    st_qa_a = S("qaa", alignment=TA_LEFT, leftIndent=22, spaceAfter=6)
    st_ref = S("ref", fontSize=13, leading=17, alignment=TA_LEFT, leftIndent=26, firstLineIndent=-26, spaceAfter=6)
    st_tcell = lambda fs, al, bold=False: S("tc", fontName="Body-Bold" if bold else "Body", fontSize=fs, leading=fs * 1.28, alignment=al)
    AL = {"l": TA_LEFT, "r": TA_RIGHT, "c": TA_CENTER}

    def mk(text):
        t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        out = []
        for s, b, i, c in tokens(t):
            if c: out.append(f'<font name="Code" size="11.5">{s}</font>')
            elif b and i: out.append(f"<b><i>{s}</i></b>")
            elif b: out.append(f"<b>{s}</b>")
            elif i: out.append(f"<i>{s}</i>")
            else: out.append(s)
        return "".join(out).replace("\n", "<br/>")

    CODE_FS, CODE_LEAD, CODE_PAD = 10, 12.6, 6
    CODE_W = FW
    MAXCH = int((CODE_W - 2 * CODE_PAD - 2) / (0.602 * CODE_FS))

    def wrap_code(line):
        if len(line) <= MAXCH: return [line]
        ind = len(line) - len(line.lstrip(" "))
        out, cont = [], " " * (ind + 4)
        cur = line
        first = True
        while len(cur) > MAXCH:
            cut = cur.rfind(" ", ind + 1 if first else len(cont) + 1, MAXCH)
            if cut <= 0: cut = MAXCH
            out.append(cur[:cut].rstrip() if cut != MAXCH else cur[:cut])
            cur = cont + cur[cut:].lstrip(" ")
            first = False
        out.append(cur); return out

    class CodeBlock(Flowable):
        def __init__(self, lines, first=True, last=True):
            super().__init__(); self.lines = lines; self.first = first; self.last = last
            self.h = len(lines) * CODE_LEAD + (CODE_PAD if first else 0) + (CODE_PAD if last else 0)
        @classmethod
        def from_text(cls, text):
            ls = []
            for ln in text.rstrip("\n").split("\n"): ls.extend(wrap_code(ln.replace("\t", "    ")))
            return cls(ls)
        def wrap(self, aw, ah): self.width = CODE_W; return CODE_W, self.h
        def split(self, aw, ah):
            if self.h <= ah: return [self]
            top = CODE_PAD if self.first else 0
            n = int((ah - top) / CODE_LEAD)
            if n < 3 or n >= len(self.lines): return []
            return [CodeBlock(self.lines[:n], self.first, False), CodeBlock(self.lines[n:], False, self.last)]
        def draw(self):
            c = self.canv; h = self.h
            c.setFillColor(colors.HexColor("#F5F5F5")); c.setStrokeColor(colors.HexColor("#BDBDBD")); c.setLineWidth(0.6)
            c.rect(0, 0, CODE_W, h, stroke=0, fill=1)
            c.line(0, 0, 0, h); c.line(CODE_W, 0, CODE_W, h)
            if self.first: c.line(0, h, CODE_W, h)
            if self.last: c.line(0, 0, CODE_W, 0)
            y = h - (CODE_PAD if self.first else 0) - CODE_LEAD * 0.80
            cw = pdfmetrics.stringWidth("0", "Code", CODE_FS)
            for ln in self.lines:
                badge = ln.find("\u2705")
                txt = ln.replace("\u2705", " ")
                c.setFillColor(colors.HexColor("#1b1b1b")); c.setFont("Code", CODE_FS)
                c.drawString(CODE_PAD, y, txt)
                if badge >= 0:   # emoji glyph not in the embedded font: draw a green check badge
                    bx = CODE_PAD + badge * cw; s = CODE_FS * 0.95
                    c.setFillColor(colors.HexColor("#4CAF50")); c.rect(bx, y - 1.5, s, s, stroke=0, fill=1)
                    c.setStrokeColor(colors.white); c.setLineWidth(1.2)
                    p = c.beginPath(); p.moveTo(bx + .2 * s, y - 1.5 + .5 * s); p.lineTo(bx + .42 * s, y - 1.5 + .27 * s); p.lineTo(bx + .8 * s, y - 1.5 + .75 * s)
                    c.drawPath(p, stroke=1, fill=0); c.setStrokeColor(colors.HexColor("#BDBDBD")); c.setLineWidth(0.6)
                y -= CODE_LEAD

    class Doc(BaseDocTemplate):
        def afterFlowable(self, fl):
            if hasattr(fl, "_toc"):
                lvl, txt = fl._toc
                self.notify("TOCEntry", (lvl, txt, self.page))

    def on_page(c, d):
        c.saveState()
        c.setStrokeColor(colors.HexColor("#222222")); c.setLineWidth(1.0)
        o = 24
        c.rect(o, o, W - 2 * o, H - 2 * o)
        c.setLineWidth(0.4); c.rect(o + 3, o + 3, W - 2 * o - 6, H - 2 * o - 6)
        if d.page > 1:
            c.setFont("Body-Italic", 10.5); c.setFillColor(colors.HexColor("#333333"))
            c.drawCentredString(W / 2, H - 1.75 * cm, HEADER_TXT)
            c.setLineWidth(0.5); c.line(LM, H - 1.95 * cm, W - RM, H - 1.95 * cm)
        c.setLineWidth(0.5); c.line(LM, 1.95 * cm, W - RM, 1.95 * cm)
        c.setFont("Body", 10.5); c.setFillColor(colors.HexColor("#333333"))
        c.drawString(LM, 1.5 * cm, FOOTER_TXT)
        c.drawRightString(W - RM, 1.5 * cm, f"Page {d.page}")
        c.restoreState()

    doc = Doc(path, pagesize=A4, leftMargin=LM, rightMargin=RM, topMargin=TM, bottomMargin=BM,
              title=TITLE + " – Internship Project Demonstration", author="Vismay Vinod", subject="BCA AI/ML Internship Project Report")
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(LM, BM, FW, H - TM - BM, id="f", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=on_page)])

    toc = TableOfContents(dotsMinLevel=0)
    toc.levelStyles = [
        S("t0", fontName="Body-Bold", fontSize=12.5, leading=16, leftIndent=0, spaceBefore=5, rightIndent=20, firstLineIndent=0),
        S("t1", fontSize=11.5, leading=14.5, leftIndent=18, rightIndent=20, firstLineIndent=0)]

    fig_n = [0]; tab_n = [0]
    story = []

    def heading(txt, style, level, center=False):
        p = Paragraph(f"<u>{mk(txt)}</u>", style)
        if level is not None: p._toc = (level, plain(txt))
        return p

    def table_flow(head, rows, widths, aligns, caption, fs):
        aligns = aligns or ["l"] * len(head)
        data = [[Paragraph(mk(h), st_tcell(fs, AL[a] if a != "r" else TA_RIGHT, True)) for h, a in zip(head, aligns)]]
        for r in rows:
            data.append([Paragraph(mk(str(x)), st_tcell(fs, AL[a])) for x, a in zip(r, aligns)])
        t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#555555")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E3E3")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                               ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
        out = []
        if caption:
            tab_n[0] += 1
            out.append(Paragraph(f"<b>Table {tab_n[0]}:</b> {mk(caption)}", S("tcap", fontSize=11.5, leading=14.5, alignment=TA_CENTER, spaceBefore=6, spaceAfter=5, keepWithNext=1)))
        out += [t, Spacer(1, 12)]
        return out

    def sig_flow(labels):
        n = len(labels); w = FW / n
        data = [[Paragraph(mk(l), S("sg", fontSize=12.5, leading=16, alignment=TA_CENTER)) for l in labels]]
        t = Table(data, colWidths=[w] * n, rowHeights=None)
        t.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.8, ink), ("TOPPADDING", (0, 0), (-1, -1), 5),
                               ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14)]))
        # draw lines only over the middle part by insetting via spacer columns is overkill; keep full-width rule per cell
        return [Spacer(1, 38), t, Spacer(1, 14)]

    def cover_flow():
        out = [Spacer(1, 1.2 * cm), Paragraph("<b>AI/ML INTERNSHIP PROJECT DEMONSTRATION</b>", S("c1", fontSize=24, leading=30, alignment=TA_CENTER)),
               Spacer(1, 22), Paragraph("Project Title", S("c2", fontSize=13, alignment=TA_CENTER, textColor=colors.HexColor("#444444"))),
               Spacer(1, 4), Paragraph(f"<b>{TITLE}</b>", S("c3", fontSize=22, leading=28, alignment=TA_CENTER)), Spacer(1, 26)]
        data = [[Paragraph(f"<b>{k}</b>", S("ck", fontSize=13.5, leading=17)), Paragraph(mk(v), S("cv", fontSize=13.5, leading=17))] for k, v in COVER_ROWS]
        t = Table(data, colWidths=[6.0 * cm, FW - 6.0 * cm])
        t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")), ("TOPPADDING", (0, 0), (-1, -1), 7),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("LINEABOVE", (0, 0), (-1, 0), 0.4, colors.HexColor("#999999"))]))
        out += [t, Spacer(1, 22), Paragraph("Notebook: AI_ML_Salary_Prediction.ipynb   |   Dataset: ai_ml_job_analysis.csv", S("c4", fontSize=11.5, alignment=TA_CENTER)),
                Paragraph(f"Repository: {REPO_URL}", S("c5", fontSize=11.5, alignment=TA_CENTER))]
        return out

    for bi, blk in enumerate(B):
        k = blk[0]
        if k == "cover": story += cover_flow()
        elif k == "h1":
            _, t, new, center = blk
            if new: story.append(PageBreak())
            else: story.append(Spacer(1, 18))
            story.append(heading(t, st_h1c if center else st_h1, 0))
        elif k == "h2":
            _, t, new = blk
            if new: story.append(PageBreak())
            story.append(heading(t, st_h2, 1))
        elif k == "h3":
            if blk[2]: story.append(PageBreak())
            nxt = B[bi + 1] if bi + 1 < len(B) else None
            if nxt is not None and nxt[0] == "img":
                wc, hc = img_size_cm(nxt[1][0], nxt[3])
                story.append(CondPageBreak(min(hc * cm + 3.2 * cm, 0.9 * (H - TM - BM))))
            story.append(heading(blk[1], st_h3, None))
        elif k == "p": story.append(Paragraph(mk(blk[1]), st_body_l if is_long(blk[1]) else st_body))
        elif k == "note": story.append(Paragraph(mk(blk[1]), st_note))
        elif k == "bul": story += [Paragraph(mk(i), st_bul_l if is_long(i) else st_bul, bulletText="•") for i in blk[1]] + [Spacer(1, 4)]
        elif k == "num": story += [Paragraph(mk(i), st_bul, bulletText=f"{n}.") for n, i in enumerate(blk[1], 1)] + [Spacer(1, 4)]
        elif k == "code": story += [CodeBlock.from_text(blk[1]), Spacer(1, 10)]
        elif k == "sp": story.append(Spacer(1, blk[1]))
        elif k == "line": story.append(Paragraph(mk(("**%s**" % blk[1]) if blk[2] else blk[1]), S("ln", fontSize=blk[3] or 14, alignment=TA_CENTER if blk[4] else TA_LEFT, spaceAfter=blk[5], leading=(blk[3] or 14) * 1.3)))
        elif k == "sig": story += sig_flow(blk[1])
        elif k == "table": story += table_flow(*blk[1:])
        elif k == "qa": story += [Paragraph(f"Q{blk[1]}. {mk(blk[2])}", st_qa_q), Paragraph(mk(blk[3]), st_qa_a)]
        elif k == "ref": story.append(Paragraph(f"[{blk[1]}]&nbsp;&nbsp;{mk(blk[2])}", st_ref))
        elif k == "toc":
            story += [PageBreak(), Paragraph("<u>TABLE OF CONTENTS</u>", st_h1c), toc]
        elif k == "img":
            _, files, caption, wcm = blk
            fig_n[0] += 1
            items = []
            for f in files:
                wc, hc = img_size_cm(f, wcm)
                items += [Image(f, width=wc * cm, height=hc * cm), Spacer(1, 5)]
            items.append(Paragraph(f"<b>Figure {fig_n[0]}:</b> {mk(caption)}", st_cap))
            for it in items[:-1]:
                if isinstance(it, Image): it.hAlign = "CENTER"
            story.append(KeepTogether(items))
    doc.multiBuild(story)
    entries = [(e[0], e[1], e[2]) for e in toc._lastEntries]
    return entries

# =============================================================================================
#                                           DOCX
# =============================================================================================
def build_docx(path, toc_entries):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    BODY, MONO = "Times New Roman", "Courier New"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3); sec.top_margin = Cm(2.7); sec.bottom_margin = Cm(2.5)
    sec.header_distance = Cm(1.2); sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True
    TEXTW = 21.0 - 4.6

    def set_font(style_or_run, name, size=None, bold=None, italic=None, underline=None, color="000000"):
        f = style_or_run.font
        f.name = name
        rpr = (style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else style_or_run.element.rPr)
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), name)
        if size is not None: f.size = Pt(size)
        if bold is not None: f.bold = bold
        if italic is not None: f.italic = italic
        if underline is not None: f.underline = underline
        if color: f.color.rgb = RGBColor.from_string(color)

    # ---- styles
    st = doc.styles
    n = st["Normal"]; set_font(n, BODY, 14)
    n.paragraph_format.space_after = Pt(7); n.paragraph_format.space_before = Pt(0); n.paragraph_format.line_spacing = 1.12
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, sb, sa in (("Heading 1", 0, 12), ("Heading 2", 14, 8), ("Heading 3", 10, 6)):
        h = st[name]; set_font(h, BODY, 18, bold=True, italic=False, underline=True)
        pf = h.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa); pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.line_spacing = 1.05
    code_st = st.add_style("CodeLine", WD_STYLE_TYPE.PARAGRAPH); code_st.base_style = n
    set_font(code_st, MONO, 10); pf = code_st.paragraph_format
    pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = Cm(0.2); pf.right_indent = Cm(0.2); pf.widow_control = False
    for nm in ("List Bullet", "List Number"):
        s = st[nm]; set_font(s, BODY, 14); s.paragraph_format.space_after = Pt(4); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for nm, ind, bold, sz in (("TOC 1", 0, True, 12.5), ("TOC 2", 0.6, False, 11.5)):
        s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH); s.element.find(qn("w:name")).set(qn("w:val"), nm.lower()); s.base_style = n; set_font(s, BODY, sz, bold=bold)
        s.paragraph_format.left_indent = Cm(ind); s.paragraph_format.space_after = Pt(2 if nm == "TOC 2" else 3)
        s.paragraph_format.space_before = Pt(4 if nm == "TOC 1" else 0); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.line_spacing = 1.0
        s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXTW), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # ---- page border (all pages incl. cover)
    sectPr = sec._sectPr
    pgb = OxmlElement("w:pgBorders"); pgb.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "24"); e.set(qn("w:color"), "222222")
        pgb.append(e)
    sectPr.find(qn("w:pgMar")).addnext(pgb)

    # ---- helpers
    def add_field(par, instr, size=10.5, text="1"):
        for kind in ("begin", None, "separate", "text", "end"):
            r = par.add_run(); set_font(r, BODY, size, color="333333")
            if kind in ("begin", "separate", "end"):
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
            elif kind is None:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
            else:
                r.text = text

    def para_border(p, bottom=False, top=False, color="555555", sz="4"):
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr")
        for side in (("top",) if top else ()) + (("bottom",) if bottom else ()):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:space"), "2"); e.set(qn("w:color"), color); b.append(e)
        pPr.append(b)

    def shade(el_pr, fill):
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); el_pr.append(sh)

    # header / footer
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(HEADER_TXT); set_font(r, BODY, 10.5, italic=True, color="333333"); para_border(hp, bottom=True)
    hp.paragraph_format.space_after = Pt(0)
    for footer in (sec.footer, sec.first_page_footer):
        fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.LEFT; fp.paragraph_format.space_after = Pt(0)
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXTW), WD_TAB_ALIGNMENT.RIGHT)
        para_border(fp, top=True)
        r = fp.add_run(FOOTER_TXT + "\tPage "); set_font(r, BODY, 10.5, color="333333")
        add_field(fp, "PAGE", 10.5)
    sec.first_page_header.paragraphs[0].text = ""

    def runs(par, text, size=None, base_bold=False, base_italic=False, color="000000"):
        for s, b, i, c in tokens(text):
            r = par.add_run(s)
            if c: set_font(r, MONO, 11.5, bold=b or base_bold, color=color)
            else: set_font(r, BODY, size, bold=(b or base_bold), italic=(i or base_italic), color=color)

    def cell_text(cell, text, fs, align, bold=False):
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.line_spacing = 1.0
        runs(p, text, fs, base_bold=bold)

    def set_cell_width(cell, cm_):
        cell.width = Cm(cm_)
        tcPr = cell._tc.get_or_add_tcPr()

    def cell_borders(cell, **sides):
        tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            v = sides.get(side)
            if v: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v)); e.set(qn("w:color"), "222222")
            else: e.set(qn("w:val"), "nil")
            b.append(e)
        tcPr.append(b)

    def fix_table_layout(tbl, widths):
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for row in tbl.rows:
            for c, w in zip(row.cells, widths): c.width = Cm(w)
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))

    fig_n = [0]; tab_n = [0]; toc_i = [0]

    def heading(txt, level, new=False, center=False, before=None):
        p = doc.add_paragraph(style=f"Heading {level}")
        if before is not None: p.paragraph_format.space_before = Pt(before)
        p.add_run(txt)
        if new: p.paragraph_format.page_break_before = True
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    # ---- block renderers
    for blk in B:
        k = blk[0]
        if k == "cover":
            def cp(text, size, bold=False, after=6, before=0, color="000000"):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
                r = p.add_run(text); set_font(r, BODY, size, bold=bold, color=color); return p
            cp("AI/ML INTERNSHIP PROJECT DEMONSTRATION", 24, True, 14, 40)
            cp("Project Title", 13, False, 2, 18, "444444")
            cp(TITLE, 22, True, 22)
            t = doc.add_table(rows=len(COVER_ROWS), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (a, b) in enumerate(COVER_ROWS):
                c0, c1 = t.rows[i].cells
                for c, txt, bd in ((c0, a, True), (c1, b, False)):
                    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(6); p.paragraph_format.space_before = Pt(6)
                    r = p.add_run(txt); set_font(r, BODY, 13.5, bold=bd)
                    cell_borders(c, top=4 if i == 0 else None, bottom=4)
            fix_table_layout(t, [6.0, TEXTW - 6.0])
            cp("Notebook: AI_ML_Salary_Prediction.ipynb   |   Dataset: ai_ml_job_analysis.csv", 11.5, False, 0, 22)
            cp(f"Repository: {REPO_URL}", 11.5, False, 0)
        elif k == "h1": heading(blk[1], 1, blk[2], blk[3], None if blk[2] else 20)
        elif k == "h2": heading(blk[1], 2, blk[2])
        elif k == "h3": heading(blk[1], 3, blk[2])
        elif k == "p":
            pp = doc.add_paragraph(); runs(pp, blk[1], 14)
            if is_long(blk[1]): pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif k == "note":
            p = doc.add_paragraph(); runs(p, blk[1], 12, base_italic=True); shade(p._p.get_or_add_pPr(), "F2F2F2")
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(12)
        elif k == "bul":
            for it in blk[1]:
                pp = doc.add_paragraph(style="List Bullet"); runs(pp, it, 14)
                if is_long(it): pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif k == "num":
            for it in blk[1]: runs(doc.add_paragraph(style="List Number"), it, 14)
        elif k == "code":
            lines = blk[1].rstrip("\n").split("\n")
            for idx, ln in enumerate(lines):
                p = doc.add_paragraph(style="CodeLine")
                r = p.add_run(ln.replace("\t", "    ")); set_font(r, MONO, 10)
                pPr = p._p.get_or_add_pPr()
                bd = OxmlElement("w:pBdr")
                for side in ("top", "left", "bottom", "right"):
                    e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "3" if side in ("left", "right") else "1"); e.set(qn("w:color"), "BDBDBD"); bd.append(e)
                pPr.append(bd); shade(pPr, "F5F5F5")
                if idx < 2 and len(lines) > 3: p.paragraph_format.keep_with_next = True
                if idx == len(lines) - 1: p.paragraph_format.space_after = Pt(10)
        elif k == "sp":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(max(blk[1], 1))
            r = p.add_run(""); set_font(r, BODY, 2)
        elif k == "line":
            _, t, bold, size, center, after = blk
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(after)
            r = p.add_run(t); set_font(r, BODY, size or 14, bold=bold)
        elif k == "sig":
            labels = blk[1]; nn = len(labels)
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = Pt(26)
            t = doc.add_table(rows=1, cols=nn); w = TEXTW / nn
            for c, lab in zip(t.rows[0].cells, labels):
                cell_borders(c, top=6)
                for j, part in enumerate(lab.split("\n")):
                    p = c.paragraphs[0] if j == 0 else c.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(4 if j == 0 else 0)
                    r = p.add_run(part); set_font(r, BODY, 12.5, bold=(j == 0))
            fix_table_layout(t, [w] * nn)
            sp2 = doc.add_paragraph(); sp2.paragraph_format.space_after = Pt(6)
        elif k == "table":
            _, head, rows, widths, aligns, caption, fs = blk
            aligns = aligns or ["l"] * len(head)
            if caption:
                tab_n[0] += 1
                cp_ = doc.add_paragraph(); cp_.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp_.paragraph_format.keep_with_next = True
                cp_.paragraph_format.space_before = Pt(6); cp_.paragraph_format.space_after = Pt(5)
                r = cp_.add_run(f"Table {tab_n[0]}: "); set_font(r, BODY, 11.5, bold=True)
                runs(cp_, caption, 11.5)
            t = doc.add_table(rows=1 + len(rows), cols=len(head)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, h, a in zip(t.rows[0].cells, head, aligns):
                cell_text(c, h, fs, a, True); shade(c._tc.get_or_add_tcPr(), "E3E3E3")
            trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)
            for ri, row in enumerate(rows, 1):
                for c, x, a in zip(t.rows[ri].cells, row, aligns): cell_text(c, str(x), fs, a)
            for row in t.rows:
                trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            fix_table_layout(t, widths)
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6)
        elif k == "qa":
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(f"Q{blk[1]}. "); set_font(r, BODY, 14, bold=True); runs(p, blk[2], 14, base_bold=True)
            a = doc.add_paragraph(); a.paragraph_format.left_indent = Cm(0.8); a.paragraph_format.space_after = Pt(6); a.alignment = WD_ALIGN_PARAGRAPH.LEFT; runs(a, blk[3], 14)
        elif k == "ref":
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(f"[{blk[1]}]\t"); set_font(r, BODY, 13); runs(p, blk[2], 13)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        elif k == "toc":
            heading("TABLE OF CONTENTS", 1, True, True).style = st["Heading 1"]
            # TOC heading must not appear in the TOC itself: use a plain paragraph with the same look instead
            doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
            hp_ = doc.add_paragraph(); hp_.alignment = WD_ALIGN_PARAGRAPH.CENTER; hp_.paragraph_format.page_break_before = True
            hp_.paragraph_format.space_after = Pt(14); hp_.paragraph_format.keep_with_next = True
            r = hp_.add_run("TABLE OF CONTENTS"); set_font(r, BODY, 18, bold=True, underline=True)
            first = True
            for ent_i, (lvl, text, pg) in enumerate(toc_entries):
                p = doc.add_paragraph(style="toc 1" if lvl == 0 else "toc 2")
                if first:
                    for kind in ("begin", "instr", "separate"):
                        rr = p.add_run()
                        if kind == "instr":
                            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-2" \\h \\z \\u '; rr._r.append(it)
                        else:
                            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); rr._r.append(fc)
                    first = False
                rr = p.add_run(f"{text}\t{pg}"); set_font(rr, BODY, 12.5 if lvl == 0 else 11.5, bold=(lvl == 0))
                if ent_i == len(toc_entries) - 1:
                    re_ = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); re_._r.append(fc)
        elif k == "img":
            _, files, caption, wcm = blk
            fig_n[0] += 1
            for j, f in enumerate(files):
                wc, hc = img_size_cm(f, wcm)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.0
                p.add_run().add_picture(f, width=Cm(wc))
            cp_ = doc.add_paragraph(); cp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp_.paragraph_format.space_before = Pt(2); cp_.paragraph_format.space_after = Pt(12)
            r = cp_.add_run(f"Figure {fig_n[0]}: "); set_font(r, BODY, 11.5, bold=True); runs(cp_, caption, 11.5)

    # ---- put manually-added OOXML children into schema order (Word is strict about this)
    ORDER = {
      "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
      "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark".split(),
      "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden".split(),
      "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook".split(),
    }
    W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    def reorder(root):
        for tag, order in ORDER.items():
            for el in root.iter(W_NS + tag):
                kids = list(el)
                def key(k):
                    nm = k.tag.replace(W_NS, "")
                    return order.index(nm) if nm in order else len(order)
                srt = sorted(kids, key=key)
                if srt != kids:
                    for k in kids: el.remove(k)
                    for k in srt: el.append(k)
    reorder(doc.element)
    for sct in doc.sections:
        for part in (sct.header, sct.footer, sct.first_page_header, sct.first_page_footer):
            reorder(part._element)
    # refresh fields (TOC page numbers) when opened in Word
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    anchor = None
    for nm in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping", "decimalSymbol", "listSeparator"):
        anchor = settings.find(W_NS + nm)
        if anchor is not None: break
    if anchor is not None: anchor.addprevious(uf)
    else: settings.append(uf)
    cpps = doc.core_properties
    cpps.title = TITLE + " – Internship Project Demonstration"; cpps.author = "Vismay Vinod"; cpps.subject = "BCA AI/ML Internship Project Report"
    doc.save(path)

if __name__ == "__main__":
    pdf = os.path.join(OUTDIR, BASENAME + ".pdf"); docx = os.path.join(OUTDIR, BASENAME + ".docx")
    entries = build_pdf(pdf)
    print("PDF built; TOC entries:", len(entries))
    build_docx(docx, entries)
    print("DOCX built")
